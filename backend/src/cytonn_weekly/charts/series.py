"""Reading, cleaning and checking one chart's series from an analyst workbook.  No drawing.

Everything here is a function of a ``Workbook`` (``weekly/workbook.py``) or of plain numbers.  It
stores nothing, draws nothing and reaches no screen.

How a series is read (``read_series``):

* The sheet is found by name and the series by the header text printed above its dates; the
  values are in the next column.  Never by a row number, because rows are added every week.  A
  missing sheet or header, a header found twice, or a number format that contradicts the spec's
  unit stops the read with one blocking defect, and nothing is taken from anywhere else.
* The table starts at the first date under the header and ends at the last number; dates that
  run on below it with no value are ignored, and one empty row inside it does not end it (the
  real dividend-yield column has one).  A run of ``GAP_ROWS`` empty rows ends the search.
* Every value is paired with the date in its own row of its own date column, and the points are
  ordered by date.  Newest-first and oldest-first tables both read.
* Nothing is repaired.  An empty or non-numeric value cell, a value with no date, a repeated
  date, or a date out of order is listed by cell reference and left out of the series and of the
  average.  Where two neighbouring dates are not in order the workbook does not say which is
  wrong, so both rows are left out.  These defects do not stop the chart.
* A chart is not drawn (``drawn`` is False) when it has a blocking defect: too few points, a newest
  point older than the spec allows (stale), or any point dated after the spec's ceiling.

Numbers are exact: a cell becomes a ``Decimal`` from its shortest text form, a fraction shown as a
percentage is multiplied by 100 once, and the average is an exact mean rounded half up once.
"""

from __future__ import annotations

import difflib
import math
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, localcontext
from typing import Any, Mapping, Optional

from openpyxl.utils import get_column_letter

from cytonn_weekly.charts.specs import AVERAGE_FROM, ChartSpec
from cytonn_weekly.checkers.digital_payments import Flag, SourceValue, check_figure, parse_number
from cytonn_weekly.common.formatting import round_half_up
from cytonn_weekly.weekly.workbook import Cell, Sheet, Workbook, WorkbookError, average_range

HEADER_ROWS = 30   # the header label is looked for in the sheet's first rows (the real one is on row 6)
GAP_ROWS = 10      # this many empty rows in a row end the table; fewer do not
LISTED = 8         # cell references named in a defect's text; the full list stays in ``cells``


@dataclass(frozen=True)
class Point:
    day: date
    value: Decimal      # in the chart's unit: a percentage is in percentage points
    row: int
    cell: str           # the value cell, "B8"


@dataclass(frozen=True)
class Defect:
    kind: str
    message: str
    cells: tuple[str, ...] = ()
    blocking: bool = False    # a blocking defect means the chart is not drawn


@dataclass(frozen=True)
class Stated:
    """The workbook's own AVERAGE cell for the series."""

    cell: str
    value: Optional[Decimal]          # in the chart's unit; None when the cell holds no saved value
    formula_range: Optional[str]      # "B8:B938", when the cell is =AVERAGE(...) over one column


@dataclass(frozen=True)
class ChartSeries:
    spec: ChartSpec
    week_ending: date
    col: str                          # the value column's letter
    points: tuple[Point, ...]         # oldest first
    defects: tuple[Defect, ...]
    stated: Optional[Stated]

    @property
    def drawn(self) -> bool:
        return not any(d.blocking for d in self.defects)

    @property
    def latest(self) -> Optional[Point]:
        return self.points[-1] if self.points else None


@dataclass(frozen=True)
class Average:
    value: Decimal      # the exact mean
    display: Decimal    # at the spec's display precision, rounded half up once
    count: int
    first: date
    last: date
    cells: str          # "B8:B935", the rows the points came from


@dataclass(frozen=True)
class AverageCheck:
    flag: Optional[Flag]    # SOURCES_DISAGREE when the workbook's AVERAGE cell prints a different figure
    note: Optional[str]     # the two ranges, when they differ


# ---------------------------------------------------------------------------
# Cells
# ---------------------------------------------------------------------------

def _filled(v: Any) -> bool:
    return v is not None and not (isinstance(v, str) and not v.strip())


def _is_number(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def _day(v: Any) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    return v if isinstance(v, date) else None


def _refs(refs: list[str]) -> str:
    return ", ".join(refs[:LISTED]) + (f" and {len(refs) - LISTED} more" if len(refs) > LISTED else "")


def _defect(kind: str, text: str, refs: list[str], blocking: bool = False) -> Defect:
    return Defect(kind, f"{text}: {_refs(refs)}" if refs else text, tuple(refs), blocking)


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------

def read_series(book: Workbook, spec: ChartSpec, week_ending: date) -> ChartSeries:
    """The chart's series, cleaned, with every defect listed.  ``week_ending`` is the report's Friday."""
    def stop(kind: str, message: str, refs: tuple[str, ...] = ()) -> ChartSeries:
        return ChartSeries(spec, week_ending, "", (), (Defect(kind, message, refs, True),), None)

    try:
        sheet = book.sheet(spec.sheet)
    except WorkbookError:
        near = difflib.get_close_matches(spec.sheet, book.sheet_names, n=3, cutoff=0.6)
        hint = f" Sheets with a similar name: {', '.join(repr(n) for n in near)}." if near else ""
        return stop("sheet_missing", f"{book.name} has no sheet {spec.sheet!r}.{hint} Nothing was read from any other sheet.")

    heads = sheet.find_all(spec.header, rows=range(1, HEADER_ROWS + 1))
    if len(heads) != 1:
        found = "no" if not heads else f"{len(heads)}"
        return stop("header_missing" if not heads else "header_ambiguous",
                    f"'{spec.sheet}' has {found} {spec.header!r} header in its first {HEADER_ROWS} rows; "
                    "the series is found by that header.", tuple(h.a1 for h in heads))
    head = heads[0]
    dcol, vcol = head.col, head.col + 1
    col = get_column_letter(vcol)

    first = next((r for r in range(head.row + 1, sheet.max_row + 1) if _day(sheet.value(r, dcol)) is not None), None)
    if first is None:
        return stop("empty", f"'{spec.sheet}' has no dated rows under {head.a1} ({spec.header!r}).")

    # The table runs to the last number, allowing for dates that go on below it with no values.
    last_active, quiet = first, 0
    for r in range(first + 1, sheet.max_row + 1):
        if _filled(sheet.value(r, dcol)) or _filled(sheet.value(r, vcol)):
            last_active, quiet = r, 0
        else:
            quiet += 1
            if quiet >= GAP_ROWS:
                break
    numeric_rows = [r for r in range(first, last_active + 1) if _is_number(sheet.value(r, vcol))]
    if not numeric_rows:
        return stop("empty", f"'{spec.sheet}' has no numbers in {col} under {head.a1} ({spec.header!r}).")
    last = numeric_rows[-1]

    dated: list[tuple[int, date]] = []
    blank, text, undated = [], [], []
    for r in range(head.row + 1, first):   # a number above the first date with nothing beside it has lost its date
        if _is_number(sheet.value(r, vcol)) and not _filled(sheet.value(r, dcol)):
            undated.append(f"{col}{r}")
    for r in range(first, last + 1):
        v, day, ref = sheet.value(r, vcol), _day(sheet.value(r, dcol)), f"{col}{r}"
        if day is not None:
            dated.append((r, day))
        if _is_number(v):
            if day is None:
                undated.append(ref)
        elif _filled(v):
            text.append(ref)
        else:
            blank.append(ref)

    pairs = list(zip(dated, dated[1:]))
    down, up = sum(a > b for (_, a), (_, b) in pairs), sum(a < b for (_, a), (_, b) in pairs)
    if down == up and down > 0:
        return stop("no_direction", f"The dates under {head.a1} run neither newest-first nor oldest-first.")
    newest_first = down >= up
    duplicate: set[int] = set()
    disorder: set[int] = set()
    for (ra, a), (rb, b) in pairs:
        if a == b:
            duplicate |= {ra, rb}
        elif (a > b) != newest_first:
            disorder |= {ra, rb}

    day_of = dict(dated)
    left_out = duplicate | disorder
    kept = [(r, day_of[r], sheet.value(r, vcol)) for r in range(first, last + 1)
            if r in day_of and r not in left_out and _is_number(sheet.value(r, vcol))]

    if kept:   # the number format must agree with the unit the spec says, or the figures would be 100 times out
        newest_row = max(kept, key=lambda k: k[1])[0]
        fmt = book.number_format(sheet.cell(newest_row, vcol))
        if ("%" in fmt) != (spec.unit == "pct"):
            want = "a percentage" if spec.unit == "pct" else "a plain number, not a percentage"
            return stop("unit", f"{col}{newest_row} is formatted {fmt!r}; this chart expects {want}. Nothing was converted.",
                        (f"{col}{newest_row}",))
    points = tuple(sorted((Point(day, parse_number(v) * spec.scale, r, f"{col}{r}") for r, day, v in kept),
                          key=lambda p: p.day))

    defects: list[Defect] = []
    for kind, what, refs in (
        ("blank_value", "Empty value cells inside the table (not plotted, not in the average)", blank),
        ("non_numeric", "Value cells that hold text (not plotted, not in the average)", text),
        ("missing_date", "Values with no readable date beside them (not plotted, not in the average)", undated),
    ):
        if refs:
            defects.append(_defect(kind, what, refs))
    for kind, what, rows in (
        ("duplicate_date", "Rows that share a date (both left out; the workbook does not say which is right)", duplicate),
        ("out_of_order", "Neighbouring rows whose dates are out of order (both left out; the workbook does not say which "
                         "is wrong)", disorder),
    ):
        if rows:
            defects.append(_defect(kind, what, [f"{col}{r}" for r in sorted(rows)]))

    stated, stated_defects = _read_stated(book, sheet, head, first, vcol, spec)
    defects += stated_defects

    if len(points) < spec.min_points:
        defects.append(Defect("too_few", f"Only {len(points)} usable points; a chart needs at least {spec.min_points}.", (), True))
    if points:
        newest, not_before = points[-1], week_ending + timedelta(days=spec.newest_not_before_days)
        if newest.day < not_before:
            defects.append(Defect(
                "stale", f"The newest point is dated {newest.day.isoformat()}; for the week ending {week_ending.isoformat()} it "
                         f"should be dated {not_before.isoformat()} or later. The workbook is not updated for this week.",
                (newest.cell,), True))
        ceiling = week_ending + timedelta(days=spec.ceiling_days)
        late = [p.cell for p in points if p.day > ceiling]
        if late:
            defects.append(_defect(
                "future", f"Points dated after {ceiling.isoformat()}, the latest date allowed for the week ending "
                          f"{week_ending.isoformat()}; the rows are not removed for you", late, True))
    return ChartSeries(spec, week_ending, col, points, tuple(defects), stated)


def _read_stated(book: Workbook, sheet: Sheet, head: Cell, first: int, vcol: int,
                 spec: ChartSpec) -> tuple[Optional[Stated], list[Defect]]:
    """The workbook's own AVERAGE cell: the value beside an "AVERAGE" label between the header and the first date."""
    # An empty range would make Sheet.find_all search the whole sheet, so it is not asked.
    labels = sheet.find_all("average", rows=range(head.row + 1, first), cols=range(head.col, head.col + 1)) \
        if first > head.row + 1 else []
    if len(labels) != 1:
        return None, [Defect("stated_missing", f"'{spec.sheet}' has no single 'AVERAGE' label under {head.a1}, so the chart's "
                                               "average is not compared with the workbook's own.")]
    cell = sheet.cell(labels[0].row, vcol)
    rng = average_range(book.formula(cell))
    formula_range = f"{rng[0]}{rng[1]}:{rng[0]}{rng[2]}" if rng else None
    if cell.number is None:
        return Stated(cell.a1, None, formula_range), [Defect(
            "stated_unreadable", f"{cell.a1} holds no saved number (if it is a formula, open the workbook in Excel and save it), "
                                 "so the chart's average is not compared with the workbook's own.", (cell.a1,))]
    return Stated(cell.a1, parse_number(cell.number) * spec.scale, formula_range), []


# ---------------------------------------------------------------------------
# The window, the average and the checks
# ---------------------------------------------------------------------------

def _years_back(day: date, years: int) -> date:
    try:
        return day.replace(year=day.year - years)
    except ValueError:    # 29 February in a year that has none
        return day.replace(year=day.year - years, day=28)


def windowed(series: ChartSeries) -> tuple[Point, ...]:
    """The points the picture shows: the spec's years back from the newest point."""
    if not series.points:
        return ()
    start = _years_back(series.points[-1].day, series.spec.window_years)
    return tuple(p for p in series.points if p.day >= start)


def average(series: ChartSeries) -> Optional[Average]:
    """The exact mean of every point dated ``AVERAGE_FROM`` or later (not just the window), rounded half up once."""
    used = [p for p in series.points if p.day >= AVERAGE_FROM]
    if not used:
        return None
    with localcontext() as ctx:
        ctx.prec = 50
        mean = sum((p.value for p in used), Decimal(0)) / len(used)
    rows = [p.row for p in used]
    return Average(mean, round_half_up(mean, series.spec.decimals), len(used), used[0].day, used[-1].day,
                   f"{series.col}{min(rows)}:{series.col}{max(rows)}")


def compare_average(series: ChartSeries) -> AverageCheck:
    """The chart's average against the workbook's own AVERAGE cell, at the display precision.

    A difference is "sources disagree": a human call, not an error in either.  The typed
    "Average = x" label on the team's chart is never read.  The note says when the two cover
    different ranges, whether or not the figures agree.
    """
    avg, st = average(series), series.stated
    if avg is None or st is None or st.value is None:
        return AverageCheck(None, None)
    note = None
    if st.formula_range and st.formula_range != avg.cells:
        note = (f"The workbook's {st.cell} averages {st.formula_range}; this chart averages {avg.cells}, "
                f"the points dated {AVERAGE_FROM.isoformat()} or later.")
    flag = check_figure("chart", series.spec.id, "average", avg.display, avg.value, series.spec.decimals,
                        extra_sources=[SourceValue(f"the workbook's AVERAGE cell {st.cell}", st.value)],
                        primary_source=f"the mean of {avg.cells}")
    return AverageCheck(flag, note)


def check_quoted(series: ChartSeries, quoted: Mapping[str, Any]) -> list[Flag]:
    """The last plotted point and the average, each against the figure the draft prints for it.

    ``quoted`` maps the section's figure keys (``spec.quoted``) to what the draft prints, as
    text ("10.0x") or a number.  The comparison is the exact-match checker's: both sides to the
    display precision, rounded half up.  A figure the draft does not print is a MISSING flag.
    A chart that is not drawn has no plotted point, so it has nothing to compare.
    """
    if not series.drawn or not series.points:
        return []
    spec, last, avg = series.spec, series.points[-1], average(series)
    flags = []
    for what, key in spec.quoted:
        if what == "point":
            field, source, origin = "last point", last.value, f"{spec.sheet}!{last.cell}"
        else:
            field, source, origin = "average", avg.value if avg else None, f"the mean of {avg.cells}" if avg else "the workbook"
        flag = check_figure("chart", spec.id, f"{field} ({key})", quoted.get(key), source, spec.decimals, primary_source=origin)
        if flag is not None:
            flags.append(flag)
    return flags
