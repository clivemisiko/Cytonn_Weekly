"""Reading Cytonn's analyst workbooks: cells found by sheet name plus a label or a date, never by a fixed row.

The workbooks are maintained by hand, so rows are added every day and blocks move.  A
reader here names the sheet (verbatim, trailing spaces included), then finds what it wants
by the text printed beside it ("T-Bills Results", "Year Open 2026") or by the date in its
row.  Every value comes back as a ``Cell`` carrying where it was read, so the review can
show "analyst input from <workbook> <sheet>!<cell>".

Values are the ones Excel last calculated and saved (``data_only``): a formula cell with no
saved value reads as empty, and the reader says so instead of computing it.
"""

from __future__ import annotations

import io
import re
import warnings
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any, Callable, Iterator, Optional, Union

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


class WorkbookError(ValueError):
    """The workbook does not hold what was asked for (a missing sheet, label or week)."""


@dataclass(frozen=True)
class Cell:
    """One value and where it was read."""

    value: Any
    sheet: str
    row: int
    col: int

    @property
    def a1(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"

    @property
    def ref(self) -> str:
        return f"'{self.sheet}'!{self.a1}" if not self.sheet.replace("_", "").isalnum() else f"{self.sheet}!{self.a1}"

    @property
    def number(self) -> Optional[float]:
        v = self.value
        return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None

    @property
    def day(self) -> Optional[date]:
        v = self.value
        if isinstance(v, datetime):
            return v.date()
        return v if isinstance(v, date) else None


def norm(value: Any) -> str:
    """A label as compared: its words in lower case, single-spaced."""
    return " ".join(str(value).split()).lower() if isinstance(value, str) else ""


class Sheet:
    """One worksheet's saved values, held as a grid (1-based rows and columns)."""

    def __init__(self, name: str, rows: list[tuple[Any, ...]]):
        self.name = name
        self._rows = rows

    @property
    def max_row(self) -> int:
        return len(self._rows)

    def value(self, row: int, col: int) -> Any:
        if row < 1 or row > len(self._rows):
            return None
        r = self._rows[row - 1]
        return r[col - 1] if 1 <= col <= len(r) else None

    def cell(self, row: int, col: int) -> Cell:
        return Cell(self.value(row, col), self.name, row, col)

    def cells(self, rows: Optional[range] = None, cols: Optional[range] = None) -> Iterator[Cell]:
        for r in rows or range(1, self.max_row + 1):
            line = self._rows[r - 1] if r <= len(self._rows) else ()
            for c in cols or range(1, len(line) + 1):
                v = line[c - 1] if c <= len(line) else None
                if v is not None:
                    yield Cell(v, self.name, r, c)

    def find(self, match: Union[str, Callable[[str], bool]], rows: Optional[range] = None,
             cols: Optional[range] = None, what: Optional[str] = None) -> Cell:
        """The first cell whose label equals ``match`` (or satisfies it).  WorkbookError if there is none."""
        found = self.find_all(match, rows, cols)
        if not found:
            raise WorkbookError(f"'{self.name}' has no {what or (match if isinstance(match, str) else 'such label')!r}")
        return found[0]

    def find_all(self, match: Union[str, Callable[[str], bool]], rows: Optional[range] = None,
                 cols: Optional[range] = None) -> list[Cell]:
        test = (lambda s: s == norm(match)) if isinstance(match, str) else match
        return [c for c in self.cells(rows, cols) if isinstance(c.value, str) and test(norm(c.value))]

    def row_for_date(self, day: date, col: int, rows: Optional[range] = None) -> Optional[int]:
        for r in rows or range(1, self.max_row + 1):
            v = self.value(r, col)
            if isinstance(v, datetime) and v.date() == day:
                return r
        return None

    def number(self, row: int, col: int, what: str) -> Cell:
        """A numeric cell; WorkbookError naming it if it is empty or holds text (an error value included)."""
        c = self.cell(row, col)
        if c.number is None:
            shown = "is empty" if c.value is None else f"holds {c.value!r}"
            raise WorkbookError(f"{what}: {c.ref} {shown} (if it is a formula, open the workbook in Excel and save it)")
        return c


class Workbook:
    """A workbook's saved values; its formulas are read only where a reader asks (``formula``)."""

    def __init__(self, source: Union[bytes, str], name: str):
        self.name = name
        self._source = source
        self._values = self._open(data_only=True)
        self._formulas = None
        self._sheets: dict[str, Sheet] = {}

    def _open(self, data_only: bool):
        src = io.BytesIO(self._source) if isinstance(self._source, bytes) else self._source
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")   # openpyxl warns about sparklines and other extensions it skips
            try:
                return load_workbook(src, read_only=True, data_only=data_only)
            except Exception as exc:  # noqa: BLE001 - any failure here means "not a readable .xlsx"
                raise WorkbookError(f"{self.name} could not be opened as an Excel workbook: {type(exc).__name__}: {exc}")

    @property
    def sheet_names(self) -> list[str]:
        return list(self._values.sheetnames)

    def sheet(self, name: str) -> Sheet:
        if name not in self._sheets:
            if name not in self._values.sheetnames:
                raise WorkbookError(f"{self.name} has no sheet {name!r}")
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                self._sheets[name] = Sheet(name, [tuple(r) for r in self._values[name].iter_rows(values_only=True)])
        return self._sheets[name]

    def formula(self, cell: Cell) -> Optional[str]:
        """The cell's formula text ("=AVERAGE(B8:B938)"), or None if it holds a typed value."""
        if self._formulas is None:
            self._formulas = self._open(data_only=False)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            for row in self._formulas[cell.sheet].iter_rows(min_row=cell.row, max_row=cell.row, min_col=cell.col,
                                                            max_col=cell.col, values_only=True):
                return row[0] if isinstance(row[0], str) and row[0].startswith("=") else None
        return None


_RANGE = re.compile(r"^=\s*AVERAGE\(\s*\$?([A-Z]{1,3})\$?(\d+)\s*:\s*\$?([A-Z]{1,3})\$?(\d+)\s*\)\s*$", re.I)


def average_range(formula: Optional[str]) -> Optional[tuple[str, int, int]]:
    """("B", 8, 938) for "=AVERAGE(B8:B938)"; None for anything else."""
    m = _RANGE.match(formula or "")
    if not m or m.group(1).upper() != m.group(3).upper():
        return None
    return m.group(1).upper(), int(m.group(2)), int(m.group(4))
