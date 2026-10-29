"""The approved weekly report as one Word file, in the real issue's order.

Allowed only when **every** weekly section of the week is approved: the tool never
publishes anything a coordinator has not signed off, and a half-approved report is not a
report.  The file is for whoever lays the issue out and loads it into the CMS; nothing
here sends or publishes it (publishing to the CMS is still a blocked stub).

What it holds, read from Cytonn Weekly #38.2026 (cytonnreport.com issue 891):

* the six sections under the issue's own headings, in its order;
* every table as a Word table, with its source line and footnotes;
* each chart the real issue carries, where the issue puts it: the tool's own picture (6.0 in wide, with a line under it
  saying it is from the analyst workbook and not verified) for a chart it drew and the coordinator approved, otherwise
  "[Chart to add: <caption>]" for the team's own chart (``_chart`` says each case);
* text carried forward from the previous issue marked as such, and a part the coordinator
  approved as missing named as missing, so neither can pass for finished copy;
* "Executive Summary:" on page 1, after the title, as the published issue (#41.2026) prints it: each section's summary
  as one paragraph led by the section's name in bold, a section with no summary left out.

A dev-mode draft (any piece drafted by the local model) is exported with a DEV MODE
watermark and banner on every page, and ``refuse_to_send`` blocks any path that would send
it, by the same test the Word summary's email path uses.
"""

from __future__ import annotations

import re
import unicodedata
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any, Optional, Union

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.shared import Inches, Pt, RGBColor

from cytonn_weekly import paths
from cytonn_weekly.charts import store as chart_store
from cytonn_weekly.charts.block import PROVENANCE, team_chart_chosen
from cytonn_weekly.common.word_layout import (
    add_hyperlink,
    add_linked_text,
    apply_report_look,
    executive_summary,
    shade_header,
    small_line,
    title,
    title_line,
)
from cytonn_weekly.digital_payments import review_store
from cytonn_weekly.digital_payments.coordinator_review import APPROVED, CoordinatorReview
from cytonn_weekly.digital_payments.delivery import DevModeDeliveryBlocked, is_dev_mode_review
from cytonn_weekly.digital_payments.fetcher import format_table_rows
from cytonn_weekly.digital_payments.providers.base import LINK_RE
from cytonn_weekly.digital_payments.summary import (
    TABLE_COLUMNS,
    TABLE_TITLE,
    _average_row,
    _style_document,
    footnote_links,
)
from cytonn_weekly.real_estate.review_run import layout as real_estate_layout
from cytonn_weekly.report_sections import SECTIONS, SectionSpec, is_current, latest_runs
from cytonn_weekly.report_types import WEEKLY, week_ending_of
from cytonn_weekly.weekly import chart_uploads
from cytonn_weekly.weekly.summaries import section_summary
from cytonn_weekly.weekly.week import Week, day_text, ordinal

DEV_LABEL = "DEV MODE DRAFT (local model): NOT FOR PUBLICATION OR CIRCULATION"
WATERMARK = "DEV MODE DRAFT"
SUMMARY_HEADING = "Executive Summary:"
CARRIED_NOTE = "[Carried forward from the previous issue: edit before publishing]"
CHART_WIDTH = Inches(6.0)
CHART_NOT_VERIFIED = f"{PROVENANCE[0].upper()}{PROVENANCE[1:]}."   # its own paragraph, so the editor can delete it
PICTURE_UNVERIFIED = "[Chart not drawn: the stored picture could not be verified]"
UPLOAD_UNREADABLE = "[Chart image not used: the uploaded picture could not be read. Upload it again.]"

# The issue's own section headings (#38.2026).
HEADINGS = {"company_updates": "Company Updates", "fixed_income": "Fixed Income", "equities": "Equities",
            "real_estate": "Real Estate", "digital_payments": "Digital Payments Weekly Highlights", "focus": "Focus of the Week"}

_RED = RGBColor(0xC0, 0x00, 0x00)
_INK = RGBColor(0x1F, 0x2A, 0x44)


class ExportNotAllowed(ValueError):
    """The weekly report cannot be exported yet; the message names each section that is not approved."""


# ---------------------------------------------------------------------------
# Which reviews make the week's report
# ---------------------------------------------------------------------------

def weekly_reviews(db_path: Optional[Union[Path, str]], period: str, today: Optional[date] = None
                   ) -> list[tuple[SectionSpec, Optional[CoordinatorReview]]]:
    """Each weekly section with its newest review of this period, or None if it has none (or an old unlabelled one)."""
    runs = latest_runs(db_path, WEEKLY, period)
    out = []
    for spec in SECTIONS:
        run = runs.get(spec.slug)
        review = None
        if run is not None and is_current(WEEKLY, period, run["run_date"], today):
            review = review_store.load_review(run["run_id"], db_path=db_path)
        out.append((spec, review))
    return out


def _state(review: Optional[CoordinatorReview]) -> str:
    if review is None:
        return "not_drafted"
    return review.decision or "in_review"


def export_status(db_path: Optional[Union[Path, str]], period: str, today: Optional[date] = None) -> dict[str, Any]:
    """Whether the week's report can be exported, and each section's state."""
    reviews = weekly_reviews(db_path, period, today)
    sections = [{"slug": spec.slug, "title": HEADINGS.get(spec.slug, spec.title), "state": _state(review),
                 "run_id": review.run_id if review else None} for spec, review in reviews]
    dev = [HEADINGS.get(spec.slug, spec.title) for spec, review in reviews if review is not None and is_dev_mode_review(review)]
    return {"period": period, "ready": all(s["state"] == APPROVED for s in sections), "sections": sections,
            "is_dev_draft": bool(dev), "dev_sections": dev, "dev_mode_label": DEV_LABEL if dev else None,
            "can_send": not dev, "file_name": file_name(period, focus_title(reviews, typed_first=True))}


def require_all_approved(reviews: list[tuple[SectionSpec, Optional[CoordinatorReview]]]) -> list[tuple[SectionSpec, CoordinatorReview]]:
    wording = {"not_drafted": "not drafted", "in_review": "still in review", "rejected": "rejected"}
    missing = [f"{HEADINGS.get(spec.slug, spec.title)} ({wording[_state(review)]})" for spec, review in reviews
               if _state(review) != APPROVED]
    if missing:
        raise ExportNotAllowed("The weekly report cannot be exported yet, because every section must be approved first. "
                               "Not approved yet: " + "; ".join(missing) + ". Approve those sections, then export again.")
    return [(spec, review) for spec, review in reviews if review is not None]


def refuse_to_send(reviews: list[tuple[SectionSpec, CoordinatorReview]]) -> None:
    """Raise if any section was drafted by the local (dev-only) model: such a report never reaches a recipient."""
    dev = [HEADINGS.get(spec.slug, spec.title) for spec, review in reviews if is_dev_mode_review(review)]
    if dev:
        raise DevModeDeliveryBlocked("This weekly report is a DEV MODE draft (local model) and must not be sent or "
                                     f"published: {', '.join(dev)} hold pieces drafted by the local model.")


# ---------------------------------------------------------------------------
# Document pieces
# ---------------------------------------------------------------------------

_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _plain(text: str) -> str:
    return _BOLD.sub(r"\1", text).replace("*", "")


def _add_text(doc, body_md: str) -> None:
    """Paragraphs and "- " bullets of a block's text; inline links become numbered sources listed after it.

    In the paragraphs each source's phrase is a real link (its "[n]" marker removed, or the marker linked when the
    phrase is not just before it); in the Sources list each address is a link."""
    text, sources = footnote_links(body_md or "")
    links = {n: (_plain(anchor), url) for n, anchor, url in sources}

    def para(line: str, style: Optional[str] = None) -> None:
        add_linked_text(doc.add_paragraph(style=style), _plain(line), links)

    for block in re.split(r"\n\s*\n", text.strip()):
        lines = [line for line in block.splitlines() if line.strip()]
        if lines and all(line.lstrip().startswith("- ") for line in lines):
            for line in lines:
                para(line.lstrip()[2:], "List Bullet")
        elif lines:
            head = [line for line in lines if not line.lstrip().startswith("- ")]
            para(" ".join(line.strip() for line in head))
            for line in lines:
                if line.lstrip().startswith("- "):
                    para(line.lstrip()[2:], "List Bullet")
    if sources:
        p = small_line(doc, "Sources:")
        for n, anchor, url in sources:
            p.add_run().add_break()
            p.add_run(f"[{n}] {anchor}: ").font.size = Pt(9)
            add_hyperlink(p, url, url).font.size = Pt(9)


def _add_grid(doc, grid: list[list[str]], header_rows: int = 1, numeric_from: int = 1) -> None:
    table = doc.add_table(rows=0, cols=len(grid[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, values in enumerate(grid):
        cells = table.add_row().cells
        for c, value in enumerate(values):
            cells[c].text = ""
            para = cells[c].paragraphs[0]
            run = para.add_run("" if value is None else str(value))
            run.font.size, run.bold = Pt(8), r < header_rows
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT if c >= numeric_from and r >= header_rows else WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_after = Pt(0)
            if r < header_rows:
                shade_header(cells[c])


def _source_line(doc, source: dict[str, Any]) -> None:
    if source and source.get("name"):
        small_line(doc, f"Source: {source['name']}")


def _add_table_block(doc, block: dict[str, Any]) -> None:
    title_line(doc, block["title"])
    if block.get("export_grid"):
        _add_grid(doc, block["export_grid"], header_rows=2)
    else:
        columns = block["columns"]
        numeric = next((n for n, c in enumerate(columns) if c.get("fmt", "text") != "text"), len(columns))
        grid = [[c["label"] for c in columns]] + [[row.get(c["key"], "") for c in columns] for row in block.get("rows", [])]
        _add_grid(doc, grid, numeric_from=numeric)
    for note in block.get("footnotes") or []:
        small_line(doc, note)
    for note in block.get("notes") or []:
        small_line(doc, note, italic=True)
    _source_line(doc, block.get("source") or {})
    failed = [r for r in block.get("source_rows", []) if r.get("error")]
    for r in failed:
        small_line(doc, f"{r.get(block['key_field'])}: {r['error']}", color=_RED, italic=True)


def _mark(doc, text: str):
    """A review mark's run, in a left-aligned paragraph: the body is justified, which would stretch a mark's words."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p.add_run(text)


def _placeholder(doc, caption: str) -> None:
    r = _mark(doc, f"[Chart to add: {caption}]")
    r.italic, r.font.color.rgb = True, _INK


def _red_line(doc, text: str) -> None:
    r = _mark(doc, text)
    r.italic, r.font.color.rgb = True, _RED


def _uploaded_chart(doc, caption: str, upload: tuple[bytes, str]) -> None:
    """The analyst's uploaded image: its lead-in, then the picture no wider than the text, keeping its proportions."""
    data, lead_in = upload
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        shape = para.add_run().add_picture(BytesIO(data))
    except Exception:  # noqa: BLE001 - a stored file python-docx cannot read is shown as a review line, never a crash
        para._p.getparent().remove(para._p)
        _placeholder(doc, caption)
        return _red_line(doc, UPLOAD_UNREADABLE)
    page = doc.sections[-1]
    text_width = page.page_width - page.left_margin - page.right_margin
    if shape.width > text_width:
        shape.height, shape.width = int(shape.height * text_width / shape.width), text_width
    shape._inline.docPr.set("descr", caption)
    para.insert_paragraph_before(lead_in).paragraph_format.keep_with_next = True


def _uploaded_charts(period: str, inputs_root: Optional[Union[Path, str]]) -> dict[str, tuple[bytes, str]]:
    """Caption to (image, lead-in) for each chart slot of the week with an uploaded image."""
    friday = week_ending_of(period)
    if friday is None:
        return {}
    week, out = Week(friday), {}
    for slot in chart_uploads.slots():
        stored = chart_uploads.image(week, slot["key"], inputs_root)
        if stored:
            out[slot["caption"]] = (stored[0], chart_uploads.lead_in(week, slot["key"], inputs_root) or slot["caption"])
    return out


def _chart(doc, caption: str, chart: Optional[dict[str, Any]], review: CoordinatorReview,
           charts_root: Optional[Union[Path, str]], upload: Optional[tuple[bytes, str]] = None) -> None:
    """One chart's place in the file.  ``chart`` is the section's chart block for this caption, if the tool has one.

    * The tool's chart was drawn, is still the one in use and its stored picture verifies (the same SHA-256 check as the
      review screen's route): the picture, 6.0 in wide, with the caption as its alt text and, in its own paragraph, the
      line saying where its history comes from.
    * The coordinator chose the team's own chart, or no workbook was uploaded: the placeholder and nothing else.
    * A workbook was uploaded but the chart was not drawn (the coordinator accepted it as missing): the placeholder and,
      in red, why.
    * The picture is drawn but its file is changed or missing now: the placeholder and a red line, never a picture
      that cannot be verified.
    * No chart block for the caption (every chart outside this feature): the placeholder, as before.
    * An image the analyst uploaded for the slot wins over all of these: it is the team's own chart.
    """
    if upload is not None:
        return _uploaded_chart(doc, caption, upload)
    if chart is None or chart["status"] == "no_workbook" or team_chart_chosen(review, chart["chart_id"]):
        return _placeholder(doc, caption)
    if chart["status"] != "drawn":
        _placeholder(doc, caption)
        why = next((d["message"] for d in chart.get("defects", []) if d.get("blocking")), "the chart could not be drawn")
        return _red_line(doc, f"[Chart not drawn: {why}]")
    try:
        png = chart_store.read_png(review.run_id, chart["chart_id"], chart.get("png_sha256"), charts_root)
    except chart_store.ChartFileError:
        _placeholder(doc, caption)
        return _red_line(doc, PICTURE_UNVERIFIED)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True      # the line below stays on the picture's page
    shape = para.add_run().add_picture(BytesIO(png), width=CHART_WIDTH)
    shape._inline.docPr.set("descr", caption)
    small_line(doc, CHART_NOT_VERIFIED)


def _add_blocks(doc, review: CoordinatorReview, charts_root: Optional[Union[Path, str]] = None,
                uploads: Optional[dict[str, tuple[bytes, str]]] = None) -> None:
    section = review.section
    chart_after: dict[str, list[str]] = section.get("chart_after") or {}
    charts = {b["caption"]: b for b in section.get("blocks", []) if b.get("kind") == "chart"}

    def place(captions: list[str]) -> None:
        for caption in captions:
            _chart(doc, caption, charts.get(caption), review, charts_root, (uploads or {}).get(caption))

    placed: set[str] = set()
    for block in real_estate_layout(section, WEEKLY):
        kind = block["kind"]
        if block.get("export") is False:
            continue   # shown on the review screen only; its content is inside another block's table
        if block.get("group_heading"):
            doc.add_heading(block["group_heading"], level=2)
        if kind == "table":
            _add_table_block(doc, block)
        elif kind == "narrative":
            headline = block.get("headline") or block.get("topic") or ""
            if block.get("group"):   # a Real Estate story under its group: "a. <headline>"
                doc.add_heading(f"{block['numeral']}. {headline}", level=3)
            else:
                doc.add_heading(headline, level=2)
            _add_text(doc, block.get("body_md") or "")
        elif kind == "computed":
            doc.add_heading(block["title"], level=2)
            _add_text(doc, block["body_md"])
            names = [s["name"] for s in block.get("sources") or [] if s.get("name")]
            if names:
                small_line(doc, "Source: " + "; ".join(names))
        elif kind == "carried":
            doc.add_heading(block["title"], level=2)
            if not block.get("edits"):   # a paragraph the reviewer rewrote (sentence_edit.apply_paragraph) is no longer carried text
                note = _mark(doc, CARRIED_NOTE)
                note.bold, note.font.color.rgb = True, _RED
            _add_text(doc, block["body_md"])
        elif kind == "supplied":
            if len(section.get("blocks", [])) > 1:
                doc.add_heading(block["title"], level=2)
            _add_text(doc, block["body_md"])
        elif kind == "unavailable":
            _red_line(doc, f"[Not included: {block['title']}. {block['reason']}]")
        # A chart block prints nothing where it stands: its place is where its caption is listed (chart_after).
        captions = chart_after.get(block["id"]) or []
        place(captions)
        placed.update(captions)
    notes = section.get("chart_notes") or []
    place([c for c in notes if c not in placed] + [c for c in charts if c not in placed and c not in notes])


def _add_digital_payments(doc, section: dict[str, Any]) -> None:
    """The weekly Digital Payments shape: highlights and the stock table as items, then the outlook."""
    outlook = section.get("outlook") or {}
    for item in section.get("items", []):
        if item.get("kind") == "highlight":
            doc.add_heading(f"{item.get('numeral', '')}. {item.get('headline', '')}".strip(". "), level=2)
            _add_text(doc, item.get("body_md") or item.get("body") or "")
        elif item.get("kind") == "stock_table":
            title_line(doc, f"{item.get('numeral', '')}. {TABLE_TITLE}".strip(". "))
            rows = item.get("rows", [])
            grid = [[label for _, label, _ in TABLE_COLUMNS]] + [[r[k] for k, _, _ in TABLE_COLUMNS] for r in format_table_rows(rows)]
            average = _average_row(outlook.get("stats") or {})
            if average:
                grid.append(average)
            _add_grid(doc, grid, numeric_from=2)
            small_line(doc, "Source: Yahoo Finance. Forward P/E is based on analyst consensus forward earnings estimates.")
    text = outlook.get("text") or (outlook.get("text_md") or "").strip("*")
    if text:
        doc.add_heading("Outlook", level=2)
        add_linked_text(doc.add_paragraph(), _plain(text))


def _watermark(doc) -> None:
    """A diagonal DEV MODE watermark behind every page, and the full label at the top of every page."""
    for sec in doc.sections:
        header = sec.header
        para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        run = para.add_run(DEV_LABEL)
        run.bold, run.font.color.rgb, run.font.size = True, _RED, Pt(9)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pict = parse_xml(
            '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">'
            '<v:shape id="DevModeWatermark" o:spid="_x0000_s2049" type="#_x0000_t136" '
            'style="position:absolute;margin-left:0;margin-top:0;width:468pt;height:117pt;rotation:315;z-index:-251658752;'
            'mso-position-horizontal:center;mso-position-horizontal-relative:margin;'
            'mso-position-vertical:center;mso-position-vertical-relative:margin" '
            'o:allowincell="f" fillcolor="#c00000" stroked="f">'
            '<v:fill opacity=".25"/>'
            f'<v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="{WATERMARK}"/>'
            "</v:shape></w:pict>")
        para.add_run()._r.append(pict)


def summary_text(section: dict[str, Any]) -> str:
    """The section's summary as one plain paragraph: links keep their words only (their sources are listed under the
    section's own pieces), bullets and paragraph breaks become spaces."""
    text = LINK_RE.sub(r"\1", section_summary(section))
    return " ".join(_plain(line.strip().removeprefix("- ")) for line in text.split("\n") if line.strip())


def report_title(period: str) -> str:
    week = week_ending_of(period)
    if week is not None:
        return f"Cytonn Weekly: week ending {ordinal(week)}"
    return f"Cytonn Weekly: {period}" if period else "Cytonn Weekly"


def build_weekly_document(reviews: list[tuple[SectionSpec, CoordinatorReview]], period: str,
                          charts_root: Optional[Union[Path, str]] = None, inputs_root: Optional[Union[Path, str]] = None):
    """The python-docx Document of an approved week (the caller has already required every section approved).

    ``charts_root`` is where the reviews' chart files are kept (default: under the data directory); ``inputs_root``
    is the weekly inputs store, where the week's uploaded chart images are."""
    doc = Document()
    _style_document(doc)
    apply_report_look(doc)
    topic = focus_title(reviews)
    doc.core_properties.title = f"{topic}, Cytonn Weekly" if topic else report_title(period)
    doc.core_properties.author = "Cytonn Weekly tool"
    dev = any(is_dev_mode_review(review) for _, review in reviews)
    if dev:
        _watermark(doc)
        banner = _mark(doc, DEV_LABEL)
        banner.bold, banner.font.color.rgb = True, _RED

    title(doc, report_title(period))
    decided = max((review.decided_at or "" for _, review in reviews), default="")
    on = f" (last on {day_text(date.fromisoformat(decided[:10]))})" if decided else ""
    small_line(doc, f"Every section checked and approved by the coordinator{on}. "
                "Charts are added by hand where marked; text marked as carried forward must be edited before publishing.")
    executive_summary(doc, SUMMARY_HEADING, [(HEADINGS.get(spec.slug, spec.title), summary_text(review.section))
                                             for spec, review in reviews])

    uploads = _uploaded_charts(period, inputs_root)
    for spec, review in reviews:
        doc.add_heading(HEADINGS.get(spec.slug, spec.title), level=1)
        section = review.section
        if "blocks" in section:
            _add_blocks(doc, review, charts_root, uploads)
        else:
            _add_digital_payments(doc, section)
    return doc


# ---------------------------------------------------------------------------
# Writing the file
# ---------------------------------------------------------------------------

def _label(period: str) -> str:
    week = week_ending_of(period)
    if week is not None:
        return week.isoformat()
    return re.sub(r"[^A-Za-z0-9]+", "-", period).strip("-")


MAX_SLUG = 80


def slug(text: str) -> str:
    """Lowercase ASCII words joined by hyphens: accents stripped, apostrophes dropped (Kenya's is kenyas), at most
    MAX_SLUG characters, cut at a word boundary."""
    text = re.sub(r"['\u2018\u2019\u02bc`]", "", text or "")
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii").lower()
    out = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    if len(out) > MAX_SLUG:
        cut = out[:MAX_SLUG + 1]
        out = cut.rsplit("-", 1)[0] if "-" in cut else out[:MAX_SLUG]
    return out.strip("-")


def focus_title(reviews, typed_first: bool = False) -> str:
    """The Focus of the Week title as the Word file prints it: its piece's headline, else the coordinator's topic.

    ``typed_first``: the topic the coordinator typed (the section's ``topic``) first, then the headline; the file
    name uses this, because the headline is drafted and changes on every redraft."""
    for spec, review in reviews:
        if spec.slug == "focus" and review is not None:
            section = review.section
            piece = next((b for b in section.get("blocks", []) if b.get("kind") == "narrative"), {})
            names = [piece.get("headline"), piece.get("topic"), section.get("topic")]
            if typed_first:
                names.insert(0, section.get("topic"))
            return " ".join(next((n for n in names if n and n.strip()), "").split())
    return ""


def file_name(period: str, topic: str = "") -> str:
    """"<focus title slug>-cytonn-weekly.docx"; with no title, "cytonn-weekly-<week ending>.docx"."""
    name = slug(topic)
    if name:
        return f"{name}-cytonn-weekly.docx"
    label = slug(_label(period))
    return f"cytonn-weekly-{label}.docx" if label else "cytonn-weekly.docx"


def export_dir(root: Optional[Union[Path, str]] = None) -> Path:
    return Path(root) if root is not None else paths.DATA_DIR / "exports" / "weekly"


def write_weekly_report(db_path: Optional[Union[Path, str]], period: str, out_dir: Optional[Union[Path, str]] = None,
                        today: Optional[date] = None, charts_root: Optional[Union[Path, str]] = None,
                        inputs_root: Optional[Union[Path, str]] = None) -> Path:
    """Write the week's approved report and return its path.  ExportNotAllowed unless every section is approved."""
    reviews = require_all_approved(weekly_reviews(db_path, period, today))
    folder = export_dir(out_dir)
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / file_name(period, focus_title(reviews, typed_first=True))
    build_weekly_document(reviews, period, charts_root, inputs_root).save(str(path))
    return path
