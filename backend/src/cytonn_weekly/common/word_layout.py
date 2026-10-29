from __future__ import annotations

import re
from typing import Optional
from urllib.parse import urlparse

from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor

# The look of the published issue #41.2026 (its PDF, checked 2026-10-28): Calibri, black, justified body; headings and
# table titles bold and underlined; table headers on the green band sampled from its tables (006C5B) in white.
BLACK = RGBColor(0x00, 0x00, 0x00)
DARK_GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_FILL = "006C5B"
FONT = "Calibri"
LINK_BLUE = RGBColor(0x05, 0x63, 0xC1)
HYPERLINK_STYLE = "Hyperlink"

_MARKER = re.compile(r"\[(\d+)\]")
_BARE_URL = re.compile(r"https?://[^\s<>\"]+")
_TRAILING = ".,;:!?)]}'\""


def _plain_font(style) -> None:
    """Drop the theme font attributes, which would win over the font set on the style."""
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    style.font.name = FONT
    rfonts.set(qn("w:eastAsia"), FONT)


def apply_report_look(doc) -> None:
    """Body justified and black in Calibri, Heading 1 to 3 black, bold and underlined, a page number in every footer."""
    normal = doc.styles["Normal"]
    _plain_font(normal)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name in ("Heading 1", "Heading 2", "Heading 3"):   # Heading 3: a Real Estate story under its group
        heading = doc.styles[name]
        _plain_font(heading)
        heading.font.color.rgb, heading.font.bold, heading.font.underline = BLACK, True, True
    doc.styles["Heading 1"].font.size = Pt(16)
    hyperlink_style(doc)
    for n, section in enumerate(doc.sections):
        if n and section.footer.is_linked_to_previous:
            continue   # it shows the footer before it, which already has the number
        para =section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))


def title(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold, r.underline, r.font.size, r.font.color.rgb = True, True, Pt(16), BLACK
    return p


def title_line(doc, text: str):
    """A table's or chart's own title line."""
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.underline = True, True
    p.paragraph_format.keep_with_next = True
    return p


def small_line(doc, text: str, *, color: RGBColor = DARK_GREY, italic: bool = False, bold: bool = False):
    """A source line, footnote or note under a table or paragraph: 9 pt, left-aligned (never stretched like the body)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.size, r.font.color.rgb, r.italic, r.bold = Pt(9), color, italic, bold
    p.paragraph_format.space_after = Pt(3)
    return p


def shade_header(cell) -> None:
    """The green header band, its text white and bold."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), HEADER_FILL)
    cell._tc.get_or_add_tcPr().append(shd)
    for para in cell.paragraphs:
        for run in para.runs:
            run.bold, run.font.color.rgb = True, WHITE


def executive_summary(doc, heading: str, items: list[tuple[str, str]]) -> None:
    """``heading`` as a Heading 1, then one paragraph per (section name, plain summary text), the name in bold.

    A section whose text is empty is left out.
    """
    doc.add_heading(heading, level=1)
    for name, text in items:
        if text.strip():
            p = doc.add_paragraph()
            p.add_run(f"{name}:").bold = True
            add_linked_text(p, f" {text.strip()}")


def hyperlink_style(doc):
    """The "Hyperlink" character style (blue, underlined), added if the document lacks it."""
    try:
        return doc.styles[HYPERLINK_STYLE]
    except KeyError:
        style = doc.styles.add_style(HYPERLINK_STYLE, WD_STYLE_TYPE.CHARACTER)
        style.font.color.rgb, style.font.underline = LINK_BLUE, True
        return style


def safe_url(url) -> Optional[str]:
    """``url`` if it is an http or https address with a host and no whitespace or control character, else None."""
    if not isinstance(url, str) or not url or any(c.isspace() or ord(c) < 32 for c in url):
        return None
    try:
        parsed = urlparse(url)
        host = parsed.hostname
    except ValueError:
        return None
    return url if parsed.scheme in ("http", "https") and host else None


def add_hyperlink(paragraph, text: str, url: str):
    """Append ``text`` as a real external link to ``url`` and return its run; a refused URL gives a plain run.

    The text goes through python-docx's run and the address through the part's relationships, so both are escaped
    by the library; no XML is built from either string."""
    if safe_url(url) is None:
        return paragraph.add_run(text)
    run = paragraph.add_run(text, style=hyperlink_style(paragraph.part.document))
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    link.set(qn("w:history"), "1")
    paragraph._p.append(link)
    link.append(run._r)   # moves the run inside the link
    return run


def _with_urls(paragraph, text: str) -> list:
    """Plain runs, with each bare http(s) address a link showing itself (trailing punctuation left outside it)."""
    runs, pos = [], 0
    for m in _BARE_URL.finditer(text):
        url = m.group(0).rstrip(_TRAILING)
        if pos < m.start():
            runs.append(paragraph.add_run(text[pos:m.start()]))
        runs.append(add_hyperlink(paragraph, url, url))
        pos = m.start() + len(url)
    if pos < len(text):
        runs.append(paragraph.add_run(text[pos:]))
    return runs


def add_linked_text(paragraph, text: str, sources: Optional[dict[int, tuple[str, str]]] = None) -> list:
    """Write ``text`` into ``paragraph`` with its "[n]" markers turned into links; return the runs written.

    ``sources`` maps n to (phrase, url).  When the text just before "[n]" ends with the phrase (whitespace
    normalised, case-sensitive), the phrase becomes the link and the marker, with the space before it, is removed;
    otherwise the marker "[n]" itself is the link, so no link is lost.  A marker with no source is left as text.
    Any other bare http(s) address becomes a link showing itself.  No other wording changes."""
    sources = sources or {}
    runs, pos = [], 0
    for m in _MARKER.finditer(text):
        n = int(m.group(1))
        if n not in sources:
            continue
        phrase, url = sources[n]
        before = text[pos:m.start()]
        words = phrase.split()
        found = re.search(r"\s+".join(map(re.escape, words)) + r"\s*$", before) if words else None
        if found:
            runs += _with_urls(paragraph, before[:found.start()])
            runs.append(add_hyperlink(paragraph, before[found.start():].rstrip(), url))
        else:
            runs += _with_urls(paragraph, before)
            runs.append(add_hyperlink(paragraph, m.group(0), url))
        pos = m.end()
    runs += _with_urls(paragraph, text[pos:])
    return runs
