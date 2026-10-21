"""Invented NASI P/E and dividend-yield sheets for the chart reader tests.

The sheet has the name and layout of the equities workbook's 'NASI-PE & Dividend Yield' with
made-up figures: the header labels on the sixth row, an AVERAGE row under them, then one row per
Friday, newest first, the P/E and the yield each beside their own date column, the yield stored as
a fraction and formatted as a percentage.  ``edit`` gets the worksheet last, to plant a defect.
The real workbook is confidential and is never in the repository.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date, timedelta
from typing import Callable, Optional

from openpyxl import Workbook as Xlsx

from cytonn_weekly.weekly.workbook import Workbook
from tests.weekly_helpers import WEEK, dt

SHEET = "NASI-PE & Dividend Yield"
PE_FORMAT, YIELD_FORMAT = r"0.0\x;\(0.0\x\)", "0.0%"
FRIDAY = WEEK.ending


def pe(n: int) -> float:
    """The invented P/E ``n`` weeks before the newest."""
    return round(10.0 + n * 0.1, 1)


def yld(n: int) -> float:
    """The invented dividend yield, as the sheet stores it: a fraction."""
    return round(0.05 + n * 0.0001, 4)


def day(n: int, newest: date = FRIDAY) -> date:
    return newest - timedelta(days=7 * n)


def nasi_bytes(weeks: int = 60, newest: date = FRIDAY, shift: int = 0, gap: int = 0, tail: int = 0,
               newest_first: bool = True, label: bool = True, stated: tuple[float, float] = (10.5, 0.052),
               edit: Optional[Callable] = None) -> bytes:
    """``shift`` moves everything down; ``gap`` leaves empty rows before the first date; ``tail`` adds dates with no
    values below the last value (the real sheet has them); ``label=False`` leaves out the AVERAGE row."""
    wb = Xlsx()
    ws = wb.active
    ws.title = SHEET
    head = 6 + shift
    ws.cell(head, 1, "Price to Earnings")
    ws.cell(head, 3, "Dividend Yield")
    first = head + 1
    if label:
        ws.cell(head + 1, 1, "AVERAGE")
        ws.cell(head + 1, 3, "AVERAGE")
        ws.cell(head + 1, 2, stated[0]).number_format = PE_FORMAT
        ws.cell(head + 1, 4, stated[1]).number_format = YIELD_FORMAT
        first = head + 2
    first += gap
    for i in range(weeks):
        n = i if newest_first else weeks - 1 - i
        d = dt(day(n, newest))
        ws.cell(first + i, 1, d)
        ws.cell(first + i, 2, pe(n)).number_format = PE_FORMAT
        ws.cell(first + i, 3, d)
        ws.cell(first + i, 4, yld(n)).number_format = YIELD_FORMAT
    for k in range(tail):
        d = dt(day(weeks + k, newest))
        ws.cell(first + weeks + 1 + k, 1, d)
        ws.cell(first + weeks + 1 + k, 3, d)
    if edit:
        edit(ws)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def book(**kw) -> Workbook:
    return Workbook(nasi_bytes(**kw), "Equities workbook")


def put(**cells):
    """An ``edit`` that sets cells: ``put(B20=None, A8=dt(...))``."""
    def edit(ws):
        for ref, value in cells.items():
            ws[ref] = value
    return edit


def formats(**cells):
    """An ``edit`` that sets number formats: ``formats(D8="0.0")``."""
    def edit(ws):
        for ref, code in cells.items():
            ws[ref].number_format = code
    return edit


def with_saved_values(xlsx: bytes, saved: dict[str, float]) -> bytes:
    """openpyxl writes a formula with no saved result, as a workbook never opened in Excel would; Excel writes one."""
    src, out = zipfile.ZipFile(io.BytesIO(xlsx)), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                for ref, value in saved.items():
                    xml, n = re.subn(rf'(<c r="{ref}"[^>]*><f>[^<]*</f>)<v></v>', rf"\g<1><v>{value}</v>", xml)
                    assert n == 1, f"no formula cell {ref} to give a saved value"
                data = xml.encode("utf-8")
            dst.writestr(item, data)
    return out.getvalue()
