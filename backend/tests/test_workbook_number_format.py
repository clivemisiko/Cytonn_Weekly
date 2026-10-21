"""Workbook.number_format: the format code Excel stored for a cell (used to check a series' unit)."""

import io

from openpyxl import Workbook as Xlsx

from cytonn_weekly.weekly.workbook import Workbook


def book() -> Workbook:
    wb = Xlsx()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 0.048
    ws["A1"].number_format = "0.0%"
    ws["A2"] = 11.2
    ws["A2"].number_format = r"0.0\x;\(0.0\x\)"
    ws["A3"] = 5
    out = io.BytesIO()
    wb.save(out)
    return Workbook(out.getvalue(), "t")


def test_the_stored_format_code_is_returned():
    b = book()
    s = b.sheet("S")
    assert b.number_format(s.cell(1, 1)) == "0.0%"
    assert b.number_format(s.cell(2, 1)) == r"0.0\x;\(0.0\x\)"
    assert b.number_format(s.cell(3, 1)) == "General"


def test_an_empty_cell_or_one_outside_the_sheet_is_general():
    b = book()
    s = b.sheet("S")
    assert b.number_format(s.cell(2, 9)) == "General"
    assert b.number_format(s.cell(500, 1)) == "General"
