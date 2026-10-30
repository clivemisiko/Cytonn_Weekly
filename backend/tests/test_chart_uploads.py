"""Optional chart images: the store, its routes, the inputs summary and the Word export (temp folders, invented images)."""

import struct
import zlib
from datetime import timedelta

import pytest
from docx.shared import Emu
from fastapi.testclient import TestClient

from api.main import create_app
from cytonn_weekly.equities import weekly as equities
from cytonn_weekly.fixed_income import weekly as fixed_income
from cytonn_weekly.common.review import build_review, check_section
from cytonn_weekly.weekly import chart_uploads, export, inputs
from cytonn_weekly.weekly.week import Week
from cytonn_weekly.report_sections import BY_SLUG
from tests.test_weekly_export_charts import CAPTIONS, PERIOD, approved, placeholders
from tests.weekly_helpers import WEEK

JPEG = b"\xff\xd8\xff\xe0" + b"invented jpeg body"
URL = f"/api/inputs/weekly/{WEEK.ending.isoformat()}/charts"


def png(width: int = 2, height: int = 1) -> bytes:
    """A real, tiny PNG (no dpi chunk, so python-docx reads it at 72 dpi)."""
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    rows = b"".join(b"\x00" + b"\x00\x00\x00" * width for _ in range(height))
    return (chart_uploads.MAGIC["png"] + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b""))


@pytest.fixture
def client(tmp_path):
    return TestClient(create_app(db_path=tmp_path / "app.db", load_dotenv=False, inputs_root=tmp_path / "weekly"))


def folder(tmp_path):
    return inputs.week_dir(WEEK, tmp_path / "weekly") / "charts"


def exported(review, charts, tmp_path):
    return export.build_weekly_document([(BY_SLUG["equities"], review)], PERIOD, charts_root=charts,
                                        inputs_root=tmp_path / "weekly")


def by_key(body):
    return {c["key"]: c for c in body["charts"]}


# ---------------------------------------------------------------------------
# Accepted and refused
# ---------------------------------------------------------------------------

def test_a_png_and_a_jpeg_are_accepted_and_the_summary_says_so(client):
    data = png()
    r = client.put(f"{URL}/tbill_yields", content=data)
    assert r.status_code == 200
    assert (by_key(r.json())["tbill_yields"]["type"], by_key(r.json())["tbill_yields"]["size"]) == ("png", len(data))
    r = client.put(f"{URL}/nasi_pe", content=JPEG)
    assert r.status_code == 200 and by_key(r.json())["nasi_pe"]["type"] == "jpeg"
    image = client.get(f"{URL}/nasi_pe/image")
    assert image.status_code == 200 and image.content == JPEG and image.headers["content-type"] == "image/jpeg"


def test_wrong_magic_bytes_under_a_png_name_are_refused(client, tmp_path):
    r = client.put(f"{URL}/tbill_yields?filename=chart.png", content=b"%PDF-1.7 not an image",
                   headers={"Content-Type": "image/png"})
    assert r.status_code == 422 and r.json()["problems"][0]["code"] == "upload_refused"
    assert r.json()["detail"] == "This is not a PNG or JPEG image. Save the chart as a PNG or JPEG and upload that."
    assert not (folder(tmp_path) / "tbill_yields.png").exists()


def test_an_empty_file_is_refused(client):
    r = client.put(f"{URL}/tbill_yields", content=b"")
    assert r.status_code == 422 and r.json()["detail"] == "The file is empty. Choose the chart's image file again."


def test_a_file_over_5_mb_is_refused(client):
    r = client.put(f"{URL}/tbill_yields", content=chart_uploads.MAGIC["png"] + b"\x00" * chart_uploads.MAX_BYTES)
    assert r.status_code == 413 and r.json()["problems"][0]["code"] == "upload_too_big"
    assert r.json()["detail"] == "The image is over the 5 MB limit. Save a smaller copy and upload that."


def test_an_unknown_slot_and_an_invalid_week_are_refused(client):
    assert client.put(f"{URL}/real_estate_chart", content=png()).status_code == 404
    assert client.get(f"{URL}/real_estate_chart/image").status_code == 404
    thursday = client.put("/api/inputs/weekly/2026-10-01/charts/tbill_yields", content=png())
    assert thursday.status_code == 422 and thursday.json()["problems"][0]["code"] == "week_not_valid"
    assert client.put("/api/inputs/weekly/not-a-date/charts/tbill_yields", content=png()).status_code == 422


# ---------------------------------------------------------------------------
# Replace, delete, the name on disk, the lead-in
# ---------------------------------------------------------------------------

def test_a_new_upload_replaces_the_old_one(client, tmp_path):
    client.put(f"{URL}/nasi_pe", content=png())
    r = client.put(f"{URL}/nasi_pe", content=JPEG)
    assert by_key(r.json())["nasi_pe"]["type"] == "jpeg"
    assert sorted(p.name for p in folder(tmp_path).iterdir()) == ["charts.json", "nasi_pe.jpg"]
    assert client.get(f"{URL}/nasi_pe/image").content == JPEG


def test_delete_removes_the_image_and_keeps_the_lead_in(client, tmp_path):
    client.put(f"{URL}/interbank", content=png())
    client.put(f"{URL}/interbank/lead-in", json={"lead_in": "Invented lead-in."})
    r = client.delete(f"{URL}/interbank")
    assert r.status_code == 200
    assert (by_key(r.json())["interbank"]["has_image"], by_key(r.json())["interbank"]["lead_in"]) == (False, "Invented lead-in.")
    assert not (folder(tmp_path) / "interbank.png").exists()
    assert client.get(f"{URL}/interbank/image").status_code == 404
    assert client.delete(f"{URL}/interbank").status_code == 404


def test_the_file_on_disk_is_named_by_the_server_not_the_client(client, tmp_path):
    client.put(f"{URL}/money_market?filename=..%2F..%2Fevil%20name.png", content=png())
    assert sorted(p.name for p in folder(tmp_path).iterdir()) == ["charts.json", "money_market.png"]
    assert not any("evil" in p.name for p in (tmp_path / "weekly").rglob("*"))


def test_the_lead_in_is_saved_cleared_and_limited(client):
    r = client.put(f"{URL}/import_cover/lead-in", json={"lead_in": "  The chart below shows   import cover. "})
    assert by_key(r.json())["import_cover"]["lead_in"] == "The chart below shows import cover."
    assert by_key(client.put(f"{URL}/import_cover/lead-in", json={"lead_in": ""}).json())["import_cover"]["lead_in"] == ""
    long = client.put(f"{URL}/import_cover/lead-in", json={"lead_in": "x" * (chart_uploads.MAX_LEAD_IN + 1)})
    assert long.status_code == 422 and long.json()["problems"][0]["code"] == "note_refused"


def test_the_summary_lists_every_slot_with_its_state(client):
    client.put(f"{URL}/nasi_yield", content=png())
    charts = client.get(f"/api/inputs/weekly/{WEEK.ending.isoformat()}").json()["charts"]
    assert [c["key"] for c in charts] == list(chart_uploads.KEYS)
    assert [c["section"] for c in charts].count("fixed_income") == 6 and [c["section"] for c in charts].count("equities") == 2
    assert [c["default_lead_in"] for c in charts] == [chart_uploads.DEFAULT_LEAD_INS[k] for k in chart_uploads.KEYS]
    assert not {c["default_lead_in"] for c in charts} & set(fixed_income.CHARTS + equities.CHARTS)
    assert [c["key"] for c in charts if c["has_image"]] == ["nasi_yield"]
    empty = by_key({"charts": charts})["tbill_91_growth"]
    assert (empty["type"], empty["size"], empty["lead_in"]) == (None, None, "")


def test_every_slot_has_a_clean_default_sentence():
    assert set(chart_uploads.DEFAULT_LEAD_INS) == set(chart_uploads.KEYS)
    for key in chart_uploads.KEYS:
        text = chart_uploads.DEFAULT_LEAD_INS[key]
        assert text.strip() and text.endswith(".")
        assert "(" not in text and ")" not in text and chr(0x2014) not in text


def test_the_export_status_counts_the_charts_with_no_image(client, tmp_path):
    assert client.get("/api/export/weekly", params={"period": PERIOD}).json()["charts_missing"] == 8
    client.put(f"{URL}/nasi_pe", content=png())
    assert client.get("/api/export/weekly", params={"period": PERIOD}).json()["charts_missing"] == 7
    # No week named and no reviews: the latest Friday's uploads count, so 8 only while that Friday's folder is empty.
    root = tmp_path / "weekly"
    assert export.export_status(tmp_path / "app.db", "", today=WEEK.ending + timedelta(days=8), inputs_root=root)["charts_missing"] == 8
    assert export.export_status(tmp_path / "app.db", "", today=WEEK.ending + timedelta(days=1), inputs_root=root)["charts_missing"] == 7


# ---------------------------------------------------------------------------
# The Word export
# ---------------------------------------------------------------------------

def test_a_filled_slot_prints_its_lead_in_and_image_and_an_empty_one_keeps_the_placeholder(tmp_path):
    review, charts = approved(tmp_path, workbook=False)
    chart_uploads.store(WEEK, "nasi_pe", png(), tmp_path / "weekly")
    chart_uploads.set_lead_in(WEEK, "nasi_pe", "Invented lead-in for the P/E chart.", tmp_path / "weekly")
    doc = exported(review, charts, tmp_path)
    texts = [p.text for p in doc.paragraphs]
    assert len(doc.inline_shapes) == 1
    assert doc.inline_shapes[0]._inline.docPr.get("descr") == "Invented lead-in for the P/E chart."
    at = next(i for i, p in enumerate(doc.paragraphs) if p._p.xpath(".//w:drawing"))
    assert texts[at - 1] == "Invented lead-in for the P/E chart."
    assert placeholders(doc) == [f"[Chart to add: {CAPTIONS[1]}]"]


def test_an_upload_wins_over_a_drawn_chart_and_a_blank_lead_in_prints_the_default_sentence(tmp_path):
    review, charts = approved(tmp_path)
    data = png(3, 1)
    chart_uploads.store(WEEK, "nasi_yield", data, tmp_path / "weekly")
    doc = exported(review, charts, tmp_path)
    texts = [p.text for p in doc.paragraphs]
    shapes = list(doc.inline_shapes)
    assert len(shapes) == 2                       # the drawn P/E chart and the uploaded dividend yield image
    uploaded = shapes[1]
    assert doc.part.related_parts[uploaded._inline.graphic.graphicData.pic.blipFill.blip.embed].blob == data
    at = [i for i, p in enumerate(doc.paragraphs) if p._p.xpath(".//w:drawing")][1]
    assert texts[at - 1] == "The chart below shows the NASI dividend yield."
    assert CAPTIONS[1] not in texts and uploaded._inline.docPr.get("descr") == texts[at - 1]
    assert "From the analyst workbook, not verified." not in texts[at:at + 2]


def test_a_wide_image_is_scaled_to_the_text_width_keeping_its_proportions(tmp_path):
    review, charts = approved(tmp_path, workbook=False)
    chart_uploads.store(WEEK, "nasi_pe", png(2000, 500), tmp_path / "weekly")   # 27.8 in wide at 72 dpi
    doc = exported(review, charts, tmp_path)
    page = doc.sections[-1]
    text_width = page.page_width - page.left_margin - page.right_margin
    [shape] = doc.inline_shapes
    assert shape.width == text_width and abs(shape.height - Emu(text_width // 4)) <= 1


def test_a_corrupt_stored_image_falls_back_to_the_placeholder_and_a_warning_line(tmp_path):
    review, charts = approved(tmp_path, workbook=False)
    chart_uploads.store(WEEK, "nasi_pe", chart_uploads.MAGIC["png"] + b"not a picture", tmp_path / "weekly")
    doc = exported(review, charts, tmp_path)
    assert len(doc.inline_shapes) == 0
    assert placeholders(doc) == [f"[Chart to add: {c}]" for c in CAPTIONS]
    [line] = [p for p in doc.paragraphs if p.text == export.UPLOAD_UNREADABLE]
    assert line.runs[0].italic and line.runs[0].font.color.rgb == export._RED


# ---------------------------------------------------------------------------
# The week the export reads uploads for
# ---------------------------------------------------------------------------

OTHER = Week(WEEK.ending - timedelta(days=7))
FI_SLOTS = [s["key"] for s in chart_uploads.slots() if s["section"] == "fixed_income"]


def fixed_income_review(week_ending: str, period: str = ""):
    """An approved Fixed Income-like review: a tbills block and the section's own chart placement, as the builder stores it."""
    content = {"section": "fixed_income", "title": "Fixed Income", "week_ending": week_ending,
               "blocks": [{"kind": "supplied", "id": "tbills", "title": "T-Bills", "body_md": "Invented T-bill text.",
                           "supplied_by": "coordinator"}],
               "chart_after": fixed_income.CHART_AFTER, "chart_notes": list(fixed_income.CHARTS)}
    review = build_review(content, check_section(content))
    for item in review.review_items:
        if item.resolution is None:
            item.resolve("accept")
    review.decide("approved")
    review.report_type, review.period = "weekly", period
    return review


def upload_fixed_income(week, root, keys=None):
    for key in keys or [k for k in FI_SLOTS if k != "interbank"]:
        chart_uploads.store(week, key, png(), root)


def test_an_unlabelled_period_exports_the_uploads_of_the_reviews_week(tmp_path):
    root = tmp_path / "weekly"
    upload_fixed_income(WEEK, root)
    doc = export.build_weekly_document([(BY_SLUG["fixed_income"], fixed_income_review(WEEK.ending.isoformat()))], "",
                                       inputs_root=root)
    texts = [p.text for p in doc.paragraphs]
    assert len(doc.inline_shapes) == 5
    uploaded = [c for k, c in zip(chart_uploads.KEYS, fixed_income.CHARTS) if k != "interbank"]
    assert not any(t == f"[Chart to add: {c}]" for t in texts for c in uploaded)
    # The three T-bill charts follow the tbills block, each under its own default sentence.
    at = texts.index("Invented T-bill text.")
    assert texts[at + 1:at + 7:2] == [chart_uploads.DEFAULT_LEAD_INS[k] for k in FI_SLOTS[:3]]
    assert all(doc.paragraphs[at + i]._p.xpath(".//w:drawing") for i in (2, 4, 6))


def test_a_fixed_income_slot_with_no_image_keeps_its_placeholder(tmp_path):
    root = tmp_path / "weekly"
    upload_fixed_income(WEEK, root)
    doc = export.build_weekly_document([(BY_SLUG["fixed_income"], fixed_income_review(WEEK.ending.isoformat()))], "",
                                       inputs_root=root)
    assert placeholders(doc) == [f"[Chart to add: {fixed_income.CHARTS[4]}]"]   # the interbank slot


def test_the_export_status_counts_uploads_for_an_unlabelled_period(tmp_path):
    root = tmp_path / "weekly"
    chart_uploads.store(WEEK, "tbill_yields", png(), root)
    status = export.export_status(tmp_path / "app.db", "", today=WEEK.ending, inputs_root=root)
    assert status["charts_missing"] == 7


def test_a_period_that_names_a_week_wins_over_the_reviews_week(tmp_path):
    root = tmp_path / "weekly"
    upload_fixed_income(OTHER, root)
    review = fixed_income_review(WEEK.ending.isoformat(), PERIOD)   # the review says WEEK, the period names OTHER
    named = f"Week ending {OTHER.ending.isoformat()}"
    assert export.export_week(named, [(BY_SLUG["fixed_income"], review)]) == OTHER
    doc = export.build_weekly_document([(BY_SLUG["fixed_income"], review)], named, inputs_root=root)
    assert len(doc.inline_shapes) == 5
    assert len(export.build_weekly_document([(BY_SLUG["fixed_income"], review)], PERIOD, inputs_root=root).inline_shapes) == 0


def test_an_equities_only_export_with_an_unlabelled_period_uses_the_equities_reviews_week(tmp_path):
    review, charts = approved(tmp_path, workbook=False)   # its section week_ending is WEEK
    review.period = ""
    root = tmp_path / "weekly"
    chart_uploads.store(WEEK, "nasi_pe", png(), root)
    reviews = [(BY_SLUG["equities"], review)]
    assert export.export_week("", reviews, today=OTHER.ending) == WEEK
    doc = export.build_weekly_document(reviews, "", charts_root=charts, inputs_root=root)
    assert len(doc.inline_shapes) == 1
    assert placeholders(doc) == [f"[Chart to add: {CAPTIONS[1]}]"]


def test_fixed_incomes_week_wins_and_a_bad_stored_week_is_skipped(tmp_path):
    review, _ = approved(tmp_path, workbook=False)   # Equities, week_ending WEEK
    fi = fixed_income_review(OTHER.ending.isoformat())
    assert export.export_week("", [(BY_SLUG["fixed_income"], fi), (BY_SLUG["equities"], review)]) == OTHER
    broken = fixed_income_review("not a date")
    assert export.export_week("", [(BY_SLUG["fixed_income"], broken), (BY_SLUG["equities"], review)]) == WEEK
    assert export.export_week("", [(BY_SLUG["fixed_income"], broken)], today=WEEK.ending + timedelta(days=3)) == WEEK
