import ast
import inspect

from docx import Document

from cytonn_weekly.common import word_layout as wl

REPORT_PACKAGES = ("cytonn_weekly.weekly", "cytonn_weekly.fixed_income", "cytonn_weekly.equities",
                   "cytonn_weekly.digital_payments", "cytonn_weekly.real_estate", "cytonn_weekly.focus",
                   "cytonn_weekly.periodic")


def test_the_shared_module_imports_no_report_type_package_and_holds_no_weekly_wording():
    source = inspect.getsource(wl)
    imported = []
    for node in ast.walk(ast.parse(source)):
        if isinstance(node, ast.Import):
            imported += [a.name for a in node.names]
        elif isinstance(node, ast.ImportFrom):
            imported.append(node.module or "")
    assert imported and all(m == "__future__" or m == "docx" or m.startswith("docx.") for m in imported), imported
    assert not any(m.startswith(REPORT_PACKAGES) for m in imported)
    assert "weekly" not in source.lower()


def test_the_executive_summary_leaves_out_a_section_with_no_summary():
    doc = Document()
    wl.executive_summary(doc, "Summary:", [("One", "First text."), ("Two", "  "), ("Three", "Third text.")])
    paras = doc.paragraphs
    assert paras[0].text == "Summary:" and paras[0].style.name == "Heading 1"
    assert [p.text for p in paras[1:]] == ["One: First text.", "Three: Third text."]
    assert paras[1].runs[0].text == "One:" and paras[1].runs[0].bold and not paras[1].runs[1].bold


def test_the_report_look_puts_a_page_number_in_every_footer():
    doc = Document()
    doc.add_section()
    wl.apply_report_look(doc)
    assert all('w:instr=" PAGE "' in s.footer._element.xml for s in doc.sections)
    assert doc.sections[0].footer._element.xml.count("w:fldSimple ") == 1   # a linked footer is not numbered twice
