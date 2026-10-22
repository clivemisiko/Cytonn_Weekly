"""The Word export's charts, state by state, read back with python-docx (invented workbooks and series only).

One section stands in for Equities: a paragraph with the id the real section's CHART_AFTER names ("valuation"), then
the two NASI chart blocks the real builder makes (charts/block.py), with their files written under the run id the way a
saved draft's are.  Each test then asks what the file holds in that state:

    drawn and approved                  the stored picture, 6.0 in wide, caption as alt text, a separate line under it
    the team's own chart chosen         the placeholder, nothing else
    no workbook uploaded                the placeholder, nothing else
    not drawn, accepted as missing      the placeholder and a red line saying why
    drawn, its file changed or missing  the placeholder and a red line, never a picture that cannot be verified
"""

import io
from datetime import timedelta

from docx import Document
from docx.shared import Inches, RGBColor

from cytonn_weekly.charts import store
from cytonn_weekly.charts.block import chart_blocks, set_team_chart
from cytonn_weekly.charts.specs import NASI_PE, NASI_YIELD
from cytonn_weekly.common.review import build_review, check_section
from cytonn_weekly.digital_payments import review_store
from cytonn_weekly.report_sections import BY_SLUG
from cytonn_weekly.weekly import export
from tests.chart_helpers import FRIDAY, book, day
from tests.test_charts_block import MISSING_BOOK, QUOTED, STATED

PERIOD = "Week ending 2026-10-02"
CAPTIONS = [NASI_PE.caption, NASI_YIELD.caption]
RED = RGBColor(0xC0, 0x00, 0x00)
VALUATION = {"kind": "supplied", "id": "valuation", "title": "Valuation", "body_md": "Invented valuation text.", "supplied_by": "coordinator"}
LOCAL_PIECE = {"kind": "narrative", "id": "highlight_x", "topic": "x", "headline": "An invented highlight", "body_md": "Invented text.",
               "claims": [], "warnings": [], "drafted_by": "local:invented-model", "search_scope": None, "preferred_domains": []}


def approved(tmp_path, *, workbook=True, team=(), extra=(), charts=True, listed=True, **kw):
    """An approved Equities-like review with its chart files written under ``tmp_path / "charts"``; returns (review, folder)."""
    kw.setdefault("stated", STATED)
    tmp_path.mkdir(parents=True, exist_ok=True)
    blocks, files = chart_blocks(book(**kw) if workbook else None, "Equities workbook", FRIDAY, QUOTED, MISSING_BOOK) if charts else ([], {})
    content = {"section": "equities", "title": "Equities", "week_ending": FRIDAY.isoformat(), "blocks": [VALUATION, *extra, *blocks],
               "chart_after": {"valuation": CAPTIONS} if listed else {}, "chart_notes": CAPTIONS if listed else []}
    review = build_review(content, check_section(content))
    for chart_id in team:
        set_team_chart(review, chart_id, True)
    for item in review.review_items:
        if item.resolution is None:
            item.resolve("accept")
    review.decide("approved")
    review.report_type, review.period = "weekly", PERIOD
    review_store.save_review(review, section="equities", db_path=tmp_path / "app.db")
    folder = tmp_path / "charts"
    store.attach(review, files)
    store.persist(review, folder)
    return review, folder


def document(review, folder):
    return export.build_weekly_document([(BY_SLUG["equities"], review)], PERIOD, charts_root=folder)


def texts(doc):
    return [p.text for p in doc.paragraphs]


def placeholders(doc):
    return [t for t in texts(doc) if t.startswith("[Chart to add: ")]


def red_lines(doc):
    return [(p.text, p.runs[0].font.color.rgb, p.runs[0].italic) for p in doc.paragraphs if p.text.startswith("[Chart not drawn")]


def embedded(doc, shape) -> bytes:
    return doc.part.related_parts[shape._inline.graphic.graphicData.pic.blipFill.blip.embed].blob


def png_path(review, folder, chart_id):
    return folder / str(review.run_id) / f"{chart_id}.png"


# ---------------------------------------------------------------------------
# A drawn chart
# ---------------------------------------------------------------------------

def test_a_drawn_chart_is_the_stored_picture_six_inches_wide_where_the_placeholder_was(tmp_path):
    review, folder = approved(tmp_path)
    doc = document(review, folder)
    shapes = list(doc.inline_shapes)
    assert len(shapes) == 2 and placeholders(doc) == [] and red_lines(doc) == []
    for shape, caption, chart_id in zip(shapes, CAPTIONS, ("nasi_pe", "nasi_yield")):
        assert shape.width == Inches(6.0) and abs(shape.height - Inches(3.0)) <= 1       # 6.0 x 3.0 in, whole
        assert shape._inline.docPr.get("descr") == caption                                # the caption is its alt text
        assert embedded(doc, shape) == png_path(review, folder, chart_id).read_bytes()    # the stored file, byte for byte
    # Where the issue puts them: straight after the valuation paragraph, each followed by its own small line (a separate
    # paragraph, so the editor can delete it), the picture itself alone in its paragraph.
    paras = doc.paragraphs
    at = [i for i, p in enumerate(paras) if p._p.xpath(".//w:drawing")]
    assert at == [texts(doc).index("Invented valuation text.") + 1, texts(doc).index("Invented valuation text.") + 3]
    for i in at:
        assert paras[i].text == "" and paras[i + 1].text == "From the analyst workbook, not verified."
        assert paras[i + 1]._p.xpath(".//w:drawing") == []


def test_the_picture_is_the_only_change_a_drawn_chart_makes_to_the_section(tmp_path):
    drawn_review, folder = approved(tmp_path)
    plain_review, plain_folder = approved(tmp_path / "plain", workbook=False)
    drawn, plain = document(drawn_review, folder), document(plain_review, plain_folder)
    without = [t for t in texts(drawn) if t != "From the analyst workbook, not verified."]
    assert [t for t in without if t] == [t for t in texts(plain) if t and not t.startswith("[Chart to add: ")]


# ---------------------------------------------------------------------------
# Every other state keeps the placeholder
# ---------------------------------------------------------------------------

def test_the_teams_own_chart_leaves_the_placeholder_and_no_red_line(tmp_path):
    review, folder = approved(tmp_path, team=("nasi_pe",))
    doc = document(review, folder)
    assert placeholders(doc) == [f"[Chart to add: {NASI_PE.caption}]"] and red_lines(doc) == []
    assert len(doc.inline_shapes) == 1 and next(iter(doc.inline_shapes))._inline.docPr.get("descr") == NASI_YIELD.caption
    both, folder2 = approved(tmp_path / "both", team=("nasi_pe", "nasi_yield"))
    doc2 = document(both, folder2)
    assert len(placeholders(doc2)) == 2 and len(doc2.inline_shapes) == 0 and red_lines(doc2) == []


def test_no_workbook_leaves_the_placeholders_and_no_red_line(tmp_path):
    review, folder = approved(tmp_path, workbook=False)
    doc = document(review, folder)
    assert placeholders(doc) == [f"[Chart to add: {c}]" for c in CAPTIONS]
    assert red_lines(doc) == [] and len(doc.inline_shapes) == 0


def test_a_chart_not_drawn_but_accepted_as_missing_keeps_the_placeholder_and_says_why_in_red(tmp_path):
    review, folder = approved(tmp_path, newest=day(1))      # the workbook was not updated for the week: stale
    doc = document(review, folder)
    assert placeholders(doc) == [f"[Chart to add: {c}]" for c in CAPTIONS] and len(doc.inline_shapes) == 0
    lines = red_lines(doc)
    assert len(lines) == 2
    for text, colour, italic in lines:
        assert text.startswith("[Chart not drawn: The newest point is dated ") and text.endswith("The workbook is not updated for this week.]")
        assert colour == RED and italic is True
    # the red line follows its placeholder
    t = texts(doc)
    assert t.index(lines[0][0]) == t.index(f"[Chart to add: {CAPTIONS[0]}]") + 1
    # choosing the team's own chart for it says nothing in red: the placeholder is what the coordinator wants
    chosen, folder2 = approved(tmp_path / "team", newest=day(1), team=("nasi_pe", "nasi_yield"))
    assert red_lines(document(chosen, folder2)) == []


def test_a_chart_dated_after_the_allowed_latest_date_is_not_drawn_and_says_so(tmp_path):
    review, folder = approved(tmp_path, newest=FRIDAY + timedelta(days=7))
    lines = red_lines(document(review, folder))
    assert len(lines) == 2 and all("dated after" in text for text, _, _ in lines)


# ---------------------------------------------------------------------------
# The picture is checked at export, as the route checks it
# ---------------------------------------------------------------------------

def test_a_changed_picture_is_not_embedded(tmp_path):
    review, folder = approved(tmp_path)
    path = png_path(review, folder, "nasi_pe")
    path.write_bytes(path.read_bytes() + b"\0")
    doc = document(review, folder)
    assert len(doc.inline_shapes) == 1 and next(iter(doc.inline_shapes))._inline.docPr.get("descr") == NASI_YIELD.caption
    assert placeholders(doc) == [f"[Chart to add: {NASI_PE.caption}]"]
    assert red_lines(doc) == [(export.PICTURE_UNVERIFIED, RED, True)]
    assert export.PICTURE_UNVERIFIED == "[Chart not drawn: the stored picture could not be verified]"


def test_a_missing_picture_is_not_embedded(tmp_path):
    review, folder = approved(tmp_path)
    png_path(review, folder, "nasi_yield").unlink()
    doc = document(review, folder)
    assert len(doc.inline_shapes) == 1 and next(iter(doc.inline_shapes))._inline.docPr.get("descr") == NASI_PE.caption
    assert placeholders(doc) == [f"[Chart to add: {NASI_YIELD.caption}]"]
    assert red_lines(doc) == [(export.PICTURE_UNVERIFIED, RED, True)]


def test_a_picture_swapped_for_another_valid_png_is_not_embedded(tmp_path):
    review, folder = approved(tmp_path)
    png_path(review, folder, "nasi_pe").write_bytes(png_path(review, folder, "nasi_yield").read_bytes())    # a real picture, the wrong one
    doc = document(review, folder)
    assert len(doc.inline_shapes) == 1 and red_lines(doc) == [(export.PICTURE_UNVERIFIED, RED, True)]


# ---------------------------------------------------------------------------
# Around the charts
# ---------------------------------------------------------------------------

def test_charts_outside_this_feature_stay_placeholders(tmp_path):
    review, folder = approved(tmp_path, charts=False)       # a section with chart_notes and no chart blocks
    doc = document(review, folder)
    assert placeholders(doc) == [f"[Chart to add: {c}]" for c in CAPTIONS] and len(doc.inline_shapes) == 0 and red_lines(doc) == []


def test_a_chart_whose_caption_is_in_no_list_is_still_placed_at_the_end_of_its_section(tmp_path):
    review, folder = approved(tmp_path, listed=False)       # no chart_after and no chart_notes: nothing says where they go
    doc = document(review, folder)
    paras = doc.paragraphs
    at = [i for i, p in enumerate(paras) if p._p.xpath(".//w:drawing")]
    assert len(doc.inline_shapes) == 2 and placeholders(doc) == [] and red_lines(doc) == []
    assert at[0] > texts(doc).index("Invented valuation text.")                                  # after the section's last block


def test_a_dev_mode_week_keeps_its_watermark_and_banner_and_embeds_the_charts_the_same_way(tmp_path):
    review, folder = approved(tmp_path, extra=(LOCAL_PIECE,))
    assert export.is_dev_mode_review(review)
    doc = document(review, folder)
    header = doc.sections[0].header
    assert header.paragraphs[0].text == export.DEV_LABEL
    assert 'string="DEV MODE DRAFT"' in header._element.xml and "rotation:315" in header._element.xml     # the watermark
    assert texts(doc)[0] == export.DEV_LABEL                                                            # the banner
    shapes = list(doc.inline_shapes)
    assert len(shapes) == 2 and all(s.width == Inches(6.0) for s in shapes) and placeholders(doc) == []
    clean, clean_folder = approved(tmp_path / "clean")
    assert not document(clean, clean_folder).sections[0].header.paragraphs[0].text      # no dev mark on a normal report


def test_the_file_written_by_the_api_route_holds_the_pictures(tmp_path):
    review, folder = approved(tmp_path)
    buf = io.BytesIO()
    document(review, folder).save(buf)
    assert len(Document(io.BytesIO(buf.getvalue())).inline_shapes) == 2                   # they survive saving and reopening
