"""The weekly Real Estate laid out as the issue prints it: numbered groups, lettered stories, REITs last.

The groups are worked out when the review is shown or exported, never stored, so the saved content and
everything read from it (review items, the summary) are the same as before.
"""

import copy
import io

from docx import Document

from api.serialize import serialize_review
from cytonn_weekly.common.review import ROMAN, grouped, letter, numbered
from cytonn_weekly.real_estate.review_run import GROUPS, REITS_GROUP, group_of, layout
from cytonn_weekly.report_sections import BY_SLUG
from cytonn_weekly.weekly import export
from cytonn_weekly.weekly.summaries import section_summary
from tests import section_helpers as sh

PERIOD = "Week ending 2026-09-25"
RESIDENTIAL, INFRASTRUCTURE, COMMERCIAL = "Residential Sector", "Infrastructure Sector", "Commercial Sector"


def _review():
    """residential, affordable_housing, infrastructure, commercial, then the REITs table (hospitality has no story)."""
    review = sh.real_estate_review()
    review.report_type, review.period = "weekly", PERIOD
    return review


def test_letters_run_a_to_z_and_on_past_z():
    assert [letter(i) for i in (0, 1, 25, 26, 27, 51, 52, 701, 702)] == ["a", "b", "z", "aa", "ab", "az", "ba", "zz", "aaa"]


def test_each_brief_maps_to_its_group():
    assert GROUPS == {"residential": RESIDENTIAL, "affordable_housing": RESIDENTIAL, "hospitality": "Hospitality Sector",
                      "infrastructure": INFRASTRUCTURE, "land": "Land Sector", "commercial": COMMERCIAL}
    assert group_of({"kind": "table", "id": "reits"}) == REITS_GROUP == "Real Estate Investment Trusts (REITs)"
    assert group_of({"kind": "narrative", "id": "subsection_x"}) is None   # no brief id: no group


def test_groups_follow_first_appearance_keep_stories_together_and_put_the_last_group_last():
    blocks = numbered([{"id": i, "g": g} for i, g in
                       [("r", "R"), ("x", "X"), ("t", "T"), ("r2", "R"), ("n", None), ("x2", "X")]])
    out = grouped(blocks, lambda b: b["g"], last="T")
    assert [b["id"] for b in out] == ["r", "r2", "x", "x2", "n", "t"]
    assert [b["numeral"] for b in out] == ["a", "b", "a", "b", "V", "a"]   # the ungrouped block keeps its numeral
    assert [b.get("group_heading") for b in out] == ["I. R", None, "II. X", None, None, "III. T"]
    assert blocks[0]["numeral"] == "I" and "group" not in blocks[0]     # the input is not changed


def test_more_groups_than_roman_numerals_and_more_than_26_stories_do_not_crash():
    out = grouped([{"id": str(i), "g": str(i)} for i in range(len(ROMAN) + 1)], lambda b: b["g"])
    assert out[-1]["group_heading"] == f"{len(ROMAN) + 1}. {len(ROMAN)}"
    many = grouped([{"id": str(i)} for i in range(30)], lambda b: "G")
    assert [b["numeral"] for b in many][-4:] == ["aa", "ab", "ac", "ad"]


def test_the_weekly_layout_numbers_groups_and_letters_stories_with_reits_last():
    review = _review()
    stored = copy.deepcopy(review.section)
    out = layout(review.section, "weekly")
    assert [(b["id"], b["numeral"], b.get("group_heading")) for b in out] == [
        ("subsection_residential", "a", f"I. {RESIDENTIAL}"),
        ("subsection_affordable_housing", "b", None),
        ("subsection_infrastructure", "a", f"II. {INFRASTRUCTURE}"),
        ("subsection_commercial", "a", f"III. {COMMERCIAL}"),
        ("reits", "a", f"IV. {REITS_GROUP}"),
    ]
    assert review.section == stored   # derived when shown: nothing saved changes
    assert section_summary(review.section) == section_summary(stored)


def test_markets_reviews_and_other_sections_are_left_as_they_are():
    review = _review()
    assert layout(review.section, "quarterly") is review.section["blocks"]
    other = {"section": "focus", "blocks": [{"id": "x", "kind": "narrative", "brief_id": "residential"}]}
    assert layout(other, "weekly") is other["blocks"]


def test_the_review_screen_gets_group_headings_and_letters():
    blocks = serialize_review(_review())["section"]["blocks"]
    assert [b["numeral"] for b in blocks] == ["a", "b", "a", "a", "a"]
    assert [b.get("group_heading") for b in blocks] == [f"I. {RESIDENTIAL}", None, f"II. {INFRASTRUCTURE}",
                                                        f"III. {COMMERCIAL}", f"IV. {REITS_GROUP}"]
    assert blocks[1]["group"] == RESIDENTIAL


def _old_review():
    """A review saved before pieces carried their brief id: its stories have no group."""
    review = _review()
    for b in review.section["blocks"]:
        b.pop("brief_id", None)
    return review


def test_a_review_without_brief_ids_still_shows_and_exports_its_stories_flat():
    review = _old_review()
    blocks = serialize_review(review)["section"]["blocks"]
    assert [b["numeral"] for b in blocks] == ["I", "II", "III", "IV", "a"]
    assert [b.get("group_heading") for b in blocks] == [None, None, None, None, f"I. {REITS_GROUP}"]
    doc = export.build_weekly_document([(BY_SLUG["real_estate"], review)], PERIOD)
    h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert h2[-1] == f"I. {REITS_GROUP}" and "Residential and mortgage lending: a development" in h2


def _paragraphs(review):
    buf = io.BytesIO()
    export.build_weekly_document([(BY_SLUG["real_estate"], review)], PERIOD).save(buf)
    return Document(io.BytesIO(buf.getvalue())).paragraphs


def test_the_export_prints_groups_as_heading_2_and_lettered_stories_as_heading_3_with_reits_last():
    paras = _paragraphs(_review())
    h2 = [p.text for p in paras if p.style.name == "Heading 2"]
    h3 = [p.text for p in paras if p.style.name == "Heading 3"]
    assert h2 == [f"I. {RESIDENTIAL}", f"II. {INFRASTRUCTURE}", f"III. {COMMERCIAL}", f"IV. {REITS_GROUP}"]
    assert h3 == ["a. Residential and mortgage lending: a development", "b. Affordable housing: a development",
                  "a. Infrastructure: a development", "a. Commercial office, retail and industrial: a development"]
    text = [p.text for p in paras]
    reits = text.index(f"IV. {REITS_GROUP}")
    assert text[reits + 1] == "REITs weekly performance"            # the table's own title line, under its group
    assert any(t.startswith("Source: ") for t in text[reits:])
    assert reits > max(text.index(h) for h in h3)


def test_heading_3_has_the_report_look():
    style = Document(io.BytesIO(_docx_bytes())).styles["Heading 3"]
    assert style.font.bold and style.font.underline and str(style.font.color.rgb) == "000000"


def _docx_bytes():
    buf = io.BytesIO()
    export.build_weekly_document([(BY_SLUG["real_estate"], _review())], PERIOD).save(buf)
    return buf.getvalue()
