"""Real hyperlinks in the Word export, and the export's file name."""

import io

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

from cytonn_weekly.common import word_layout as wl
from cytonn_weekly.report_sections import BY_SLUG
from cytonn_weekly.weekly import export
from tests import section_helpers as sh

PERIOD = "Week ending 2026-10-02"
URL = "https://example.org/story?a=1&b=2"


def _para():
    doc = Document()
    wl.apply_report_look(doc)
    return doc, doc.add_paragraph()


def test_the_helper_makes_an_external_link_with_the_hyperlink_style():
    doc, p = _para()
    run = wl.add_hyperlink(p, "the story", URL)
    assert len(p.hyperlinks) == 1 and p.hyperlinks[0].address == URL and p.hyperlinks[0].text == "the story"
    rid = p._p.find(qn("w:hyperlink")).get(qn("r:id"))
    assert p.part.rels[rid].is_external and p.part.rels[rid].reltype == RT.HYPERLINK and p.part.rels[rid].target_ref == URL
    assert run.style.name == "Hyperlink"
    style = doc.styles["Hyperlink"]
    assert style.font.underline and str(style.font.color.rgb) == "0563C1"
    saved = io.BytesIO()
    doc.save(saved)                                   # the "&" in the address is escaped by the library
    assert Document(io.BytesIO(saved.getvalue())).paragraphs[0].hyperlinks[0].address == URL


@pytest.mark.parametrize("bad", ["javascript:alert(1)", "file:///etc/passwd", "data:text/html,<b>x</b>", "http://",
                                 "https://exa mple.org/", "ftp://example.org/", "", "https://example.org/\nx"])
def test_a_refused_address_is_written_as_plain_text(bad):
    _, p = _para()
    wl.add_hyperlink(p, "text", bad)
    assert wl.safe_url(bad) is None and p.hyperlinks == [] and p.text == "text"


def test_a_found_phrase_becomes_the_link_and_its_marker_goes():
    _, p = _para()
    wl.add_linked_text(p, "It recovered public land valued at Sh40 million, including three government houses [1]. "
                          "Streamlinefeed [2] reported it.",
                       {1: ("recovered public land valued at Sh40 million, including three government houses", URL),
                        2: ("Streamlinefeed", "https://streamlinefeed.example/")})
    assert p.text == ("It recovered public land valued at Sh40 million, including three government houses. "
                      "Streamlinefeed reported it.")
    assert [(h.text, h.address) for h in p.hyperlinks] == [
        ("recovered public land valued at Sh40 million, including three government houses", URL),
        ("Streamlinefeed", "https://streamlinefeed.example/")]


def test_a_phrase_matches_across_extra_whitespace_but_keeps_the_text():
    _, p = _para()
    wl.add_linked_text(p, "Prices  rose\tsharply [1] again.", {1: ("Prices rose sharply", URL)})
    assert p.text == "Prices  rose\tsharply again." and p.hyperlinks[0].text == "Prices  rose\tsharply"


def test_a_phrase_not_found_keeps_the_marker_as_the_link():
    _, p = _para()
    wl.add_linked_text(p, "Prices rose [1].", {1: ("land values", URL)})
    assert p.text == "Prices rose [1]." and [(h.text, h.address) for h in p.hyperlinks] == [("[1]", URL)]


def test_the_match_is_case_sensitive():
    _, p = _para()
    wl.add_linked_text(p, "Growth [1] held.", {1: ("growth", URL)})
    assert p.text == "Growth [1] held." and p.hyperlinks[0].text == "[1]"


def test_a_repeated_phrase_links_the_occurrence_just_before_the_marker():
    _, p = _para()
    wl.add_linked_text(p, "growth slowed while growth [1] rose", {1: ("growth", URL)})
    assert p.text == "growth slowed while growth rose"
    assert p.runs[0].text == "growth slowed while " and [h.text for h in p.hyperlinks] == ["growth"]


def test_a_marker_with_no_source_is_left_alone():
    _, p = _para()
    wl.add_linked_text(p, "Table [3] shows it.", {1: ("Table", URL)})
    assert p.text == "Table [3] shows it." and p.hyperlinks == []


def test_a_bare_address_is_linked_without_its_trailing_punctuation():
    _, p = _para()
    text = "See https://a.example/x/y). Then (https://b.example), or http://c.example/z."
    wl.add_linked_text(p, text)
    assert p.text == text
    assert [h.address for h in p.hyperlinks] == ["https://a.example/x/y", "https://b.example", "http://c.example/z"]
    assert all(h.text == h.address for h in p.hyperlinks)


def test_the_executive_summary_links_a_bare_address():
    doc = Document()
    wl.executive_summary(doc, "Summary:", [("One", "Read https://a.example/r.")])
    p = doc.paragraphs[1]
    assert p.text == "One: Read https://a.example/r." and p.hyperlinks[0].address == "https://a.example/r"


ONE, TWO = "https://a.example/one", "https://b.example/two"


def test_a_real_estate_export_with_two_sources_has_four_links_and_two_targets():
    review = sh.real_estate_review()
    review.report_type, review.period = "weekly", PERIOD
    narratives = [b for b in review.section["blocks"] if b["kind"] == "narrative"]
    for b in narratives:
        b["body_md"] = ""
    narratives[0]["body_md"] = (f"The county [recovered public land]({ONE}) valued at Sh40 million. "
                                f"[Streamlinefeed]({TWO}) reported more.")
    buf = io.BytesIO()
    export.build_weekly_document([(BY_SLUG["real_estate"], review)], PERIOD).save(buf)
    doc = Document(io.BytesIO(buf.getvalue()))
    paras = doc.paragraphs
    assert sum(len(p._p.findall(qn("w:hyperlink"))) for p in paras) == 4       # two in the text, two in Sources
    assert {h.address for p in paras for h in p.hyperlinks} == {ONE, TWO}
    targets = [r.target_ref for r in doc.part.rels.values() if r.reltype == RT.HYPERLINK]
    assert sorted(targets) == [ONE, TWO]
    text = [p.text for p in paras]
    assert "The county recovered public land valued at Sh40 million. Streamlinefeed reported more." in text
    sources = next(p for p in paras if p.text.startswith("Sources:"))
    assert f"[1] recovered public land: {ONE}" in sources.text and f"[2] Streamlinefeed: {TWO}" in sources.text
    assert [h.text for h in sources.hyperlinks] == [ONE, TWO]


# --- the file name ------------------------------------------------------------------------------------------

@pytest.mark.parametrize("topic, expected", [
    ("Kenya's Balance of Payments Report", "kenyas-balance-of-payments-report"),
    ("Kenya’s Balance of Payments Report", "kenyas-balance-of-payments-report"),
    ("Côte d'Ivoire: Économie naïve", "cote-divoire-economie-naive"),
    ("Q3'2026 & H1/2026: 100% <growth>?", "q32026-h1-2026-100-growth"),
    ("  --Rates!!  ", "rates"),
    ("", ""),
    ("!!! ???", ""),
])
def test_slug(topic, expected):
    assert export.slug(topic) == expected


def test_a_long_topic_is_cut_at_a_word_boundary():
    full = export.slug("market " * 30)
    out = export.slug("Market " * 30)
    assert len(out) <= 80 and not out.endswith("-") and full.startswith(out) and full[len(out)] == "-"
    assert export.slug("a" * 100) == "a" * 80


def test_the_file_name_follows_the_topic_or_the_week():
    assert export.file_name(PERIOD, "Kenya's Balance of Payments Report") == \
        "kenyas-balance-of-payments-report-cytonn-weekly.docx"
    assert export.file_name(PERIOD) == "cytonn-weekly-2026-10-02.docx"
    assert export.file_name(PERIOD, "!!!") == "cytonn-weekly-2026-10-02.docx"
    assert export.file_name("") == "cytonn-weekly.docx"
    name = export.file_name(PERIOD, 'a/b\\c"d\'e')
    assert not any(c in name for c in '/\\"\'') and name == "a-b-c-de-cytonn-weekly.docx"


def test_the_focus_title_names_the_file_and_the_document():
    review = sh.focus_review(topic="SSA Eurobonds")
    review.report_type, review.period = "weekly", PERIOD
    reviews = [(BY_SLUG["focus"], review)]
    assert export.focus_title(reviews) == "SSA Eurobonds: a development"     # the headline the file prints
    doc = export.build_weekly_document(reviews, PERIOD)
    assert doc.core_properties.title == "SSA Eurobonds: a development, Cytonn Weekly"
