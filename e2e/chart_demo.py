"""Invented equities workbooks for the chart demo weeks, and the Equities draft kept off the network for them.

The charts are drawn from the equities workbook's 'NASI-PE & Dividend Yield' sheet.  The demo week of the other tests
(2026-10-09) has a small synthetic workbook that cannot draw them (too few weeks, no number formats), so four other
Fridays are used here, each a different case.  Every figure is invented (tests/weekly_helpers.py), and nothing in
backend/ is changed:

    2026-09-25  a complete workbook: both charts are drawn, nothing is flagged on them
    2026-09-18  the same, with rows the reader leaves out of the P/E series (an empty value cell, a repeated date)
    2026-09-11  a workbook not yet updated for the week: its newest row is the Friday before, so the charts are stale
    2026-09-04  no workbook at all: the charts are information, not findings

The Equities draft for these weeks gets fetchers that fail at once, so it makes no outside request and the parts that
need an outside source are unavailable.  Other weeks are untouched.
"""
import io
from datetime import date

from openpyxl import load_workbook

from cytonn_weekly.equities import weekly as equities
from cytonn_weekly.weekly import inputs
from cytonn_weekly.weekly.week import Week
from tests import weekly_helpers as wh

CLEAN, DEFECTS, STALE, NONE = (Week(date(2026, 9, 25)), Week(date(2026, 9, 18)), Week(date(2026, 9, 11)), Week(date(2026, 9, 4)))
WEEKS = {w.ending for w in (CLEAN, DEFECTS, STALE, NONE)}
SHEET = "NASI-PE & Dividend Yield"


def workbook(week: Week, *, updated: bool = True, defects: bool = False) -> bytes:
    """The synthetic equities workbook with enough NASI weeks to draw, its AVERAGE cells set to the mean of the rows the
    reader will use (so a week with no defect is clean), and, if asked, rows the reader leaves out."""
    wb = load_workbook(io.BytesIO(wh.equities_workbook_bytes(week, updated=updated)))
    ws = wb[SHEET]
    out = {11, 12, 20} if defects else set()          # rows 11 and 12 share a date; B20 is empty
    if defects:
        ws["A12"] = ws["A11"].value
        ws["B20"] = None
    for col, avg in ((2, "B7"), (4, "D7")):
        values = [ws.cell(r, col).value for r in range(8, 68) if r not in out or col == 4]
        ws[avg] = sum(values) / len(values)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def seed() -> None:
    """Store the demo workbooks under the weekly inputs of the data folder this process runs on."""
    inputs.store(CLEAN, "equities_workbook", workbook(CLEAN), "Equities (invented, complete).xlsx")
    inputs.store(DEFECTS, "equities_workbook", workbook(DEFECTS, defects=True), "Equities (invented, with defects).xlsx")
    inputs.store(STALE, "equities_workbook", workbook(STALE, updated=False), "Equities (invented, not updated).xlsx")


def offline_equities() -> None:
    """For the demo weeks only, the Equities draft reads no outside source: each fetcher fails at once."""
    real = equities.build_equities_weekly_review

    def down(*_args, **_kwargs):
        raise ConnectionError("E2E: the chart demo weeks use no outside source")

    offline = equities.Fetchers(usd_rates=down, price_list=down, read_box=down, afx_page=down, bulletin=down, previous_issue=down)

    def build(week_ending=None, **kwargs):
        if week_ending in WEEKS:
            kwargs.setdefault("fetchers", offline)
        return real(week_ending=week_ending, **kwargs)

    equities.build_equities_weekly_review = build
