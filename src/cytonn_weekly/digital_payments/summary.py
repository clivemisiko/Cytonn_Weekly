"""Word-doc summary of an approved Digital Payments review, for Edwin and Liz.

CLAUDE.md's process flow: after the coordinator's accuracy pass and before publish,
Edwin and Liz (insight recipients, not approvers) get a Word doc of the
coordinator-checked values for their own awareness.  This module produces that file
and stops there; getting it to them (email, shared drive) is not built.

Only an APPROVED review produces a document.  A missing, in-progress or rejected
review raises SummaryNotAllowed, so a draft can never be mistaken for the checked
version.  The document carries the drafted section as the coordinator approved it
(highlights, table, outlook) and none of the review mechanics: no flags,
resolutions or notes.

The section's stock-table rows are the raw fetcher rows, so they go through
format_table_rows() exactly as the review screen and the checker do, and the
figures read in the report's own format (372.7, (4.0%), 21.1x).  Highlight citation
links are rendered footnote-style: the anchor text gets a [n] marker and the URLs
are listed under the paragraph.

Typical usage
-------------
    path = write_latest_summary()          # latest saved review must be approved
    path = write_summary(review)           # or an approved CoordinatorReview you hold
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Optional, Union

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from cytonn_weekly.digital_payments import review_store
from cytonn_weekly.digital_payments.coordinator_review import APPROVED, CoordinatorReview
from cytonn_weekly.digital_payments.fetcher import format_table_rows
from cytonn_weekly.digital_payments.providers.base import LINK_RE

SUMMARY_DIR = Path(__file__).parents[3] / "data" / "summaries"

TITLE = "Digital Payments: weekly insight summary"
TABLE_TITLE = "Digital Payments NYSE and LSE Stock Performance"
TABLE_COLUMNS = [
    ("company", "Company", False), ("ticker", "Ticker", False),
    ("current_price", "Price", True), ("prior_close", "Price 7 days ago", True),
    ("ytd_open", "YTD open", True), ("wow_pct", "w/w %", True),
    ("ytd_pct", "YTD %", True), ("forward_pe", "Forward P/E", True),
]
_INK = RGBColor(0x1F, 0x2A, 0x44)
_GREY = RGBColor(0x59, 0x59, 0x59)
_RED = RGBColor(0xC0, 0x00, 0x00)


class SummaryNotAllowed(ValueError):
    """The review is missing or not approved, so no summary may be produced."""


# ---------------------------------------------------------------------------
# Guard
# ---------------------------------------------------------------------------

def require_approved(review: Optional[CoordinatorReview]) -> CoordinatorReview:
    """Return the review if (and only if) it is approved; raise SummaryNotAllowed otherwise."""
    if review is None:
        raise SummaryNotAllowed("there is no saved review to summarize")
    if review.decision != APPROVED:
        state = review.decision or "still in progress"
        raise SummaryNotAllowed(f"review run {review.run_id} is {state}, not approved; no summary produced")
    if not review.is_approvable:  # a decision field alone is not proof; every item must really be accepted
        raise SummaryNotAllowed(
            f"review run {review.run_id} is marked approved but not every item is resolved and accepted"
        )
    return review


# ---------------------------------------------------------------------------
# Content helpers
# ---------------------------------------------------------------------------

def footnote_links(body_md: str) -> tuple[str, list[tuple[int, str, str]]]:
    """Replace inline ``[anchor](url)`` links with ``anchor [n]``; return text and [(n, anchor, url)].

    One number per distinct URL, in order of first appearance.
    """
    numbers: dict[str, int] = {}
    sources: list[tuple[int, str, str]] = []

    def repl(m) -> str:
        anchor, url = m.group(1), m.group(2)
        if url not in numbers:
            numbers[url] = len(numbers) + 1
            sources.append((numbers[url], anchor, url))
        return f"{anchor} [{numbers[url]}]"

    return LINK_RE.sub(repl, body_md), sources


def _fmt_date(iso: str) -> str:
    d = datetime.fromisoformat(iso)
    return f"{d.day} {d:%b %Y}"


def _table_rows(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [r for it in items if it.get("kind") == "stock_table" for r in it.get("rows", [])]


def _average_row(stats: dict[str, Any]) -> Optional[list[str]]:
    """The table's Average row, from the outlook stats (already checked against the table)."""
    if not any(k in stats for k in ("avg_wow_pct", "avg_ytd_pct", "avg_forward_pe")):
        return None
    shown = format_table_rows([{
        "company": "Average", "ticker": "", "current_price": None, "prior_close": None, "ytd_open": None,
        "wow_pct": stats.get("avg_wow_pct"), "ytd_pct": stats.get("avg_ytd_pct"),
        "forward_pe": stats.get("avg_forward_pe"),
    }])[0]
    return [shown[k] if k in ("company", "wow_pct", "ytd_pct", "forward_pe") else "" for k, _, _ in TABLE_COLUMNS]


# ---------------------------------------------------------------------------
# Document building
# ---------------------------------------------------------------------------

def _style_document(doc) -> None:
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    heading = doc.styles["Heading 2"]
    heading.font.name, heading.font.size, heading.font.bold = "Calibri", Pt(12), True
    heading.font.color.rgb = _INK
    heading.paragraph_format.space_before, heading.paragraph_format.space_after = Pt(14), Pt(4)
    heading.paragraph_format.keep_with_next = True
    for style in (normal, heading):  # theme font attributes would override the font set above
        rfonts = style.element.rPr.rFonts
        rfonts.set(qn("w:eastAsia"), "Calibri")
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)


def _small(doc, text: str, *, color: RGBColor = _GREY, italic: bool = False, bold: bool = False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size, r.font.color.rgb, r.italic, r.bold = Pt(9), color, italic, bold
    p.paragraph_format.space_after = Pt(3)
    return p


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _add_table(doc, rows: list[dict[str, Any]], stats: dict[str, Any]) -> None:
    shown = format_table_rows(rows)
    body = [[r[k] for k, _, _ in TABLE_COLUMNS] for r in shown]
    avg = _average_row(stats)
    if avg:
        body.append(avg)

    table = doc.add_table(rows=1, cols=len(TABLE_COLUMNS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [1.55, 0.55, 0.65, 0.8, 0.7, 0.65, 0.65, 0.75]
    for i, (_, label, numeric) in enumerate(TABLE_COLUMNS):
        c = table.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(label)
        run.bold, run.font.size = True, Pt(9)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT if numeric else WD_ALIGN_PARAGRAPH.LEFT
        _shade(c, "E7EAF0")
    for ri, values in enumerate(body):
        is_avg = bool(avg) and ri == len(body) - 1
        cells = table.add_row().cells
        for i, ((_, _, numeric), v) in enumerate(zip(TABLE_COLUMNS, values)):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(v)
            run.font.size, run.bold = Pt(9), is_avg
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT if numeric else WD_ALIGN_PARAGRAPH.LEFT
    for row in table.rows:
        row.cells[0].width = Inches(widths[0])
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
        for c in row.cells:
            c.paragraphs[0].paragraph_format.space_after = Pt(0)

    failed = [f"{r['company']} ({r['ticker']})" for r in rows if r.get("error")]
    if failed:
        _small(doc, f"Price data could not be retrieved for: {', '.join(failed)}.", italic=True)
    if "advancers" in stats and "decliners" in stats:
        _small(doc, f"Week over week: {stats['advancers']} advancers, {stats['decliners']} decliners.")
    dates = sorted({r["current_price_date"] for r in rows if r.get("current_price_date")})
    asof = f", prices as of {_fmt_date(dates[-1])}" if dates else ""
    _small(doc, f"Source: Yahoo Finance{asof}. Forward P/E is based on analyst consensus forward earnings estimates.")


def build_summary_document(review: CoordinatorReview):
    """The python-docx Document for an approved review (raises SummaryNotAllowed otherwise)."""
    require_approved(review)
    sec = review.section
    doc = Document()
    _style_document(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.author = "Cytonn Weekly tool"

    drafted_by = {i.get("drafted_by") for i in sec["items"] if i.get("kind") == "highlight"}
    drafted_by.add(sec["outlook"].get("drafted_by"))
    if any(str(d).startswith("local:") for d in drafted_by):
        p = doc.add_paragraph()
        r = p.add_run("DEV MODE DRAFT (local model): NOT FOR CIRCULATION")
        r.bold, r.font.color.rgb = True, _RED

    p = doc.add_paragraph()
    r = p.add_run(TITLE)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(20), _INK
    p.paragraph_format.space_after = Pt(2)
    _small(doc, f"Report week {_fmt_date(sec['week_start'])} to {_fmt_date(sec['week_end'])}", bold=True)
    _small(doc, f"Checked and approved by the coordinator on {_fmt_date(review.decided_at)}. "
                "For your awareness ahead of publication.")

    found = sum(1 for it in sec["items"] if it.get("kind") == "highlight")
    if sec.get("shortfall"):
        _small(doc, f"Only {found} of {found + sec['shortfall']} highlights were found this week.",
               color=_RED, italic=True)

    for it in sec["items"]:
        if it["kind"] == "highlight":
            doc.add_heading(f"{it['numeral']}. {it['headline']}", level=2)
            text, sources = footnote_links(it.get("body_md") or it["body"])
            doc.add_paragraph(text)
            if sources:
                p = _small(doc, "Sources:")
                for n, anchor, url in sources:
                    p.add_run().add_break()
                    r = p.add_run(f"[{n}] {anchor}: {url}")
                    r.font.size, r.font.color.rgb = Pt(9), _GREY
        elif it["kind"] == "stock_table":
            doc.add_heading(f"{it['numeral']}. {TABLE_TITLE}", level=2)
            _add_table(doc, it["rows"], sec["outlook"].get("stats") or {})

    doc.add_heading("Outlook", level=2)
    outlook = sec["outlook"]
    p = doc.add_paragraph()
    p.add_run(outlook.get("text") or outlook["text_md"].strip("*")).italic = True
    return doc


# ---------------------------------------------------------------------------
# Writing the file
# ---------------------------------------------------------------------------

def default_summary_path(review: CoordinatorReview) -> Path:
    return SUMMARY_DIR / f"digital_payments_summary_{review.section['week_end']}_run{review.run_id}.docx"


def write_summary(review: Optional[CoordinatorReview], out_path: Optional[Union[Path, str]] = None) -> Path:
    """Write the .docx for an approved review and return its path.

    Raises SummaryNotAllowed (writing nothing) if ``review`` is None, undecided,
    rejected or not truly approved.  ``out_path`` defaults to
    data/summaries/digital_payments_summary_<week end>_run<id>.docx; an existing file
    at the path is replaced.
    """
    doc = build_summary_document(review)  # raises SummaryNotAllowed before anything is written
    path = Path(out_path) if out_path is not None else default_summary_path(review)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def write_latest_summary(
    out_path: Optional[Union[Path, str]] = None,
    *,
    section: str = review_store.SECTION,
    db_path: Optional[Union[Path, str]] = None,
) -> Path:
    """Summarize the latest saved review for ``section``, which must be approved.

    Deliberately the latest, not the latest approved one: if a newer draft exists
    that is still in progress or was rejected, this raises rather than quietly
    summarizing an older week.
    """
    return write_summary(review_store.load_latest_review(section, db_path=db_path), out_path)
